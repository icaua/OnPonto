from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections.abc import Iterable, Mapping
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.importadores.interpretacao_batidas import interpretar_batidas


NOME_ADAPTER = "xlsx_ponto_generico"
HORARIO_RE = re.compile(r"\b([01]?\d|2[0-3]):([0-5]\d)\b")


class ErroImportacaoXlsxPontoGenerico(ValueError):
    """Erro esperado e apresentável durante a leitura deste XLSX de ponto."""

    def __init__(self, codigo: str, mensagem: str) -> None:
        self.codigo = codigo
        self.mensagem = mensagem
        super().__init__(mensagem)


def normalizar_texto(valor: Any) -> str:
    if valor is None:
        return ""
    texto = str(valor).strip()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(char for char in texto if not unicodedata.combining(char))
    texto = re.sub(r"\s+", " ", texto)
    return texto.upper()


def texto_limpo(valor: Any) -> str:
    if valor is None:
        return ""
    return re.sub(r"\s+", " ", str(valor).strip())


def construir_lookup_mescladas(ws: Worksheet) -> dict[tuple[int, int], Any]:
    lookup: dict[tuple[int, int], Any] = {}
    for faixa in ws.merged_cells.ranges:
        valor = ws.cell(faixa.min_row, faixa.min_col).value
        for row in range(faixa.min_row, faixa.max_row + 1):
            for col in range(faixa.min_col, faixa.max_col + 1):
                lookup[(row, col)] = valor
    return lookup


def valor_celula(ws: Worksheet, row: int, col: int, lookup_mescladas: dict[tuple[int, int], Any]) -> Any:
    valor = ws.cell(row, col).value
    if valor is None:
        return lookup_mescladas.get((row, col))
    return valor


def valores_linha(ws: Worksheet, row: int, lookup_mescladas: dict[tuple[int, int], Any]) -> list[Any]:
    return [valor_celula(ws, row, col, lookup_mescladas) for col in range(1, ws.max_column + 1)]


def linha_eh_inicio_bloco(valores: list[Any]) -> bool:
    return any("NUMERO DE FUNCI" in normalizar_texto(valor) for valor in valores)


def detectar(workbook: Workbook) -> bool:
    """Checagem estrutural barata: alguma planilha tem um bloco marcado por
    uma célula contendo 'NUMERO DE FUNCIONÁRIO'."""

    for ws in workbook.worksheets:
        lookup_mescladas = construir_lookup_mescladas(ws)
        for row_idx in range(1, max(ws.max_row - 1, 1)):
            if linha_eh_inicio_bloco(valores_linha(ws, row_idx, lookup_mescladas)):
                return True
    return False


def _valor_funcionario(funcionario: Any, campo: str) -> Any:
    if isinstance(funcionario, Mapping):
        return funcionario.get(campo)
    return getattr(funcionario, campo, None)


def _indexar_funcionarios(
    funcionarios: Iterable[Any],
) -> tuple[dict[str, list[Any]], dict[str, list[Any]]]:
    por_codigo: dict[str, list[Any]] = {}
    por_nome: dict[str, list[Any]] = {}

    for funcionario in funcionarios:
        codigo_bruto = _valor_funcionario(funcionario, "codigo")
        codigo = "" if codigo_bruto is None else str(codigo_bruto).strip()
        nome = normalizar_texto(_valor_funcionario(funcionario, "nome"))
        if codigo:
            por_codigo.setdefault(codigo, []).append(funcionario)
        if nome:
            por_nome.setdefault(nome, []).append(funcionario)

    return por_codigo, por_nome


def _localizar_funcionario(
    codigo: str | None,
    nome: str,
    por_codigo: dict[str, list[Any]],
    por_nome: dict[str, list[Any]],
) -> Any | None:
    if codigo:
        candidatos_codigo = por_codigo.get(codigo, [])
        if len(candidatos_codigo) == 1:
            return candidatos_codigo[0]

    candidatos_nome = por_nome.get(normalizar_texto(nome), [])
    if len(candidatos_nome) == 1:
        return candidatos_nome[0]
    return None


def _id_registro(
    arquivo_nome: str,
    chave_funcionario: tuple[str, str],
    funcionario_id: Any | None,
    data_registro: str,
    batidas: list[str],
) -> str:
    conteudo = json.dumps(
        {
            "arquivo": arquivo_nome,
            "funcionario": chave_funcionario,
            "funcionario_resolvido": funcionario_id,
            "data": data_registro,
            "batidas": batidas,
        },
        ensure_ascii=False,
        separators=(",", ":"),
        sort_keys=True,
    )
    digest = hashlib.sha256(conteudo.encode("utf-8")).hexdigest()[:20]
    return f"{NOME_ADAPTER}-{digest}"


def extrair_codigo_funcionario(valores: list[Any]) -> str | None:
    for valor in valores:
        if "NUMERO DE FUNCI" not in normalizar_texto(valor):
            continue
        numeros = re.findall(r"\d+", str(valor))
        if numeros:
            return numeros[-1].lstrip("0") or "0"
    return None


def eh_rotulo_nome(valor: Any) -> bool:
    normalizado = normalizar_texto(valor)
    return bool(re.search(r"\bNOME\b", normalizado))


def eh_valor_nome(valor: Any) -> bool:
    normalizado = normalizar_texto(valor)
    if not normalizado:
        return False
    rotulos_invalidos = ["NUMERO DE FUNCI", "DEPARTAMENTO", "UNSET"]
    return not any(rotulo in normalizado for rotulo in rotulos_invalidos) and normalizado != "NOME"


def extrair_nome_funcionario(valores: list[Any], linha: int) -> str:
    for indice, valor in enumerate(valores):
        if not eh_rotulo_nome(valor):
            continue

        texto = texto_limpo(valor)
        inline = re.split(r"\bNOME\b\s*[:\-]?\s*", texto, flags=re.IGNORECASE, maxsplit=1)
        if len(inline) == 2 and eh_valor_nome(inline[1]):
            return inline[1]

        for candidato in valores[indice + 1 :]:
            if eh_valor_nome(candidato):
                return texto_limpo(candidato)

    return f"Funcionário sem nome linha {linha}"


def extrair_dia(valor: Any) -> int | None:
    if isinstance(valor, bool) or valor is None:
        return None
    if isinstance(valor, int):
        return valor if 1 <= valor <= 31 else None
    if isinstance(valor, float) and valor.is_integer():
        dia = int(valor)
        return dia if 1 <= dia <= 31 else None
    texto = texto_limpo(valor)
    if re.fullmatch(r"\d{1,2}", texto):
        dia = int(texto)
        return dia if 1 <= dia <= 31 else None
    return None


def formatar_hora(valor: time) -> str:
    return f"{valor.hour:02d}:{valor.minute:02d}"


def extrair_horarios(valor: Any) -> list[str]:
    if value_is_time(valor):
        return [formatar_hora(valor)]
    if isinstance(valor, datetime):
        return [formatar_hora(valor.time())]

    texto = str(valor or "")
    horarios = []
    for match in HORARIO_RE.finditer(texto):
        hora = int(match.group(1))
        minuto = int(match.group(2))
        horarios.append(f"{hora:02d}:{minuto:02d}")
    return horarios


def value_is_time(valor: Any) -> bool:
    return isinstance(valor, time) and not isinstance(valor, datetime)


def montar_marcacao(data_marcacao: date, horarios: list[str]) -> dict[str, Any]:
    interpretacao = interpretar_batidas(horarios)
    return {
        "data": data_marcacao.isoformat(),
        **interpretacao,
        "horarios_extraidos": horarios,
    }


def parse_xlsx_ponto_generico(
    file_path: str,
    *,
    arquivo_nome: str,
    mes: int,
    ano: int,
    funcionarios: Iterable[Any] = (),
) -> dict[str, Any]:
    """Lê um XLSX no layout 'largo' (bloco por funcionário, com uma linha de
    dias e uma linha de marcações em colunas)."""

    caminho = Path(file_path)
    workbook = load_workbook(caminho, data_only=True)

    por_codigo, por_nome = _indexar_funcionarios(funcionarios)
    nome_arquivo = str(arquivo_nome)
    registros: list[dict[str, Any]] = []
    total_linhas_validas = 0

    for ws in workbook.worksheets:
        lookup_mescladas = construir_lookup_mescladas(ws)

        for row_idx in range(1, max(ws.max_row - 1, 1)):
            valores_header = valores_linha(ws, row_idx, lookup_mescladas)
            if not linha_eh_inicio_bloco(valores_header):
                continue

            codigo = extrair_codigo_funcionario(valores_header)
            nome = extrair_nome_funcionario(valores_header, row_idx)
            linha_marcacoes = row_idx + 1
            linha_dias = row_idx + 2

            funcionario = _localizar_funcionario(codigo, nome, por_codigo, por_nome)
            funcionario_id = _valor_funcionario(funcionario, "id") if funcionario is not None else None
            encontrado = funcionario is not None and funcionario_id is not None
            nome_cadastrado = _valor_funcionario(funcionario, "nome") if encontrado else None
            chave_funcionario = ("codigo", codigo) if codigo else ("nome", normalizar_texto(nome))

            for col_idx in range(1, ws.max_column + 1):
                dia = extrair_dia(valor_celula(ws, linha_dias, col_idx, lookup_mescladas))
                if dia is None:
                    continue
                try:
                    data_marcacao = date(ano, mes, dia)
                except ValueError:
                    continue

                valor_marcacoes = valor_celula(ws, linha_marcacoes, col_idx, lookup_mescladas)
                horarios = extrair_horarios(valor_marcacoes)
                total_linhas_validas += 1

                resultado_interpretacao = interpretar_batidas(horarios)
                interpretacao = {
                    "entrada": resultado_interpretacao["entrada"],
                    "saida_intervalo": resultado_interpretacao["saida_almoco"],
                    "retorno_intervalo": resultado_interpretacao["retorno_almoco"],
                    "saida": resultado_interpretacao["saida"],
                }
                status = "nao_conferido" if resultado_interpretacao["status"] == "normal" else "conferir"
                pendencias = (
                    [resultado_interpretacao["observacoes"]] if resultado_interpretacao["observacoes"] else []
                )
                if not encontrado:
                    pendencias.append("Funcionário não cadastrado para esta empresa")

                data_iso = data_marcacao.isoformat()
                registros.append(
                    {
                        "id": _id_registro(
                            nome_arquivo,
                            chave_funcionario,
                            funcionario_id if encontrado else None,
                            data_iso,
                            horarios,
                        ),
                        "funcionario": {
                            "id": funcionario_id if encontrado else None,
                            "codigo_origem": codigo,
                            "nome_origem": nome,
                            "nome_cadastrado": nome_cadastrado,
                            "encontrado": encontrado,
                        },
                        "data": data_iso,
                        "batidas_originais": horarios,
                        "interpretacao": interpretacao,
                        "status": status,
                        "pendencias": pendencias,
                        # O mês/ano já vêm do parâmetro de competência: esta
                        # planilha não guarda o próprio ano, então nunca há
                        # como o dia estar fora da competência informada.
                        "fora_da_competencia": False,
                        "selecionado": encontrado,
                        "origem": {"tipo": NOME_ADAPTER, "arquivo": nome_arquivo},
                    }
                )

    if not total_linhas_validas:
        raise ErroImportacaoXlsxPontoGenerico(
            "arquivo_vazio", "O arquivo XLSX não contém registros de ponto."
        )

    registros.sort(key=lambda item: (normalizar_texto(item["funcionario"]["nome_origem"]), item["data"]))
    return {"total_linhas_validas": total_linhas_validas, "registros": registros}
