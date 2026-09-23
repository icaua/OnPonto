from datetime import date, time, timedelta
from html import escape
from io import BytesIO
import re
import unicodedata

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import HTMLResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.apuracao.service import apurar_competencia
from app.database.session import get_db
from app.database.models import Funcionario


router = APIRouter(prefix="/relatorios", tags=["Relatórios"])

ESTILO_IMPRESSAO = """
        :root {
          --brand-primary: #16845B;
          --brand-dark: #103C32;
          --brand-light: #DFF3EA;
          --surface-primary: #F7F9F8;
          --surface-white: #FFFFFF;
          --text-primary: #202723;
          --text-secondary: #5F6B66;
          --border-default: #E1E5E3;
          --warning: #E59F14;
          --warning-surface: #FDF3E1;
          --warning-text: #704900;
          color: var(--text-primary);
          font-family: Arial, sans-serif;
        }
        body {
          margin: 0;
          background: var(--surface-primary);
        }
        main {
          max-width: 980px;
          margin: 0 auto;
          padding: 32px;
          background: #fff;
          min-height: 100vh;
        }
        header {
          border-bottom: 2px solid var(--brand-primary);
          margin-bottom: 24px;
          padding-bottom: 16px;
        }
        h1, h2 {
          margin: 0 0 8px;
        }
        .meta {
          display: flex;
          gap: 18px;
          flex-wrap: wrap;
          color: var(--text-secondary);
          font-size: 14px;
        }
        button {
          border: 0;
          background: var(--brand-primary);
          color: #fff;
          border-radius: 6px;
          padding: 10px 14px;
          cursor: pointer;
          margin-bottom: 20px;
        }
        table {
          width: 100%;
          border-collapse: collapse;
          margin: 12px 0 28px;
          font-size: 13px;
        }
        th, td {
          border: 1px solid var(--border-default);
          padding: 8px;
          text-align: left;
        }
        th {
          background: var(--brand-light);
        }
        footer {
          border-top: 1px solid var(--border-default);
          color: var(--text-secondary);
          font-size: 12px;
          padding-top: 12px;
        }
        @media print {
          body {
            background: #fff;
          }
          main {
            padding: 0;
            max-width: none;
          }
          button {
            display: none;
          }
        }
"""

FORMATO_DATA_EXCEL = "dd/mm/yyyy"
FORMATO_HORA_EXCEL = "hh:mm"
FORMATO_DURACAO_EXCEL = "[h]:mm"

ROTULOS_SITUACAO = {
    "fora_vinculo": "Fora do vínculo",
    "conferido": "Conferido",
    "pendente": "Pendente",
    "indisponivel": "Indisponível",
}

ROTULOS_STATUS_DIA = {
    "fora_vinculo": "Fora do vínculo",
    "normal": "Normal",
    "falta": "Falta",
    "atestado": "Atestado",
    "folga": "Folga",
    "folga_compensatoria": "Folga compensatória",
    "feriado": "Feriado",
    "domingo": "Domingo",
    "sem_expediente": "Sem expediente",
    "trabalho_externo": "Trabalho externo",
    "afastamento": "Afastamento",
}


def texto_seguro_excel(valor: object) -> str | None:
    if valor is None:
        return None
    texto = str(valor)
    if not texto:
        return None
    if texto.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + texto
    return texto


def slug_nome_arquivo(valor: str) -> str:
    normalizado = unicodedata.normalize("NFKD", valor)
    ascii_seguro = normalizado.encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", ascii_seguro).strip("_").lower()
    return slug[:80] or "empresa"


def data_excel(valor: str | None) -> date | None:
    if not valor:
        return None
    try:
        return date.fromisoformat(valor)
    except ValueError:
        return None


def hora_excel(valor: str | None) -> time | None:
    if not valor:
        return None
    try:
        return time.fromisoformat(valor)
    except ValueError:
        return None


def duracao_excel(minutos: int | None) -> timedelta | None:
    return timedelta(minutes=minutos) if minutos is not None else None


def valor_calculado_excel(valor: object) -> object:
    return "Indisponível" if valor is None else valor


def saldo_excel(item: dict) -> timedelta | str | None:
    if item.get("pendente_calculo"):
        return None
    atraso = item.get("atraso_minutos")
    extra = item.get("extra_minutos")
    if atraso is None or extra is None:
        return None
    saldo_minutos = extra - atraso
    if saldo_minutos < 0:
        horas, minutos = divmod(abs(saldo_minutos), 60)
        return f"-{horas:02d}:{minutos:02d}"
    return duracao_excel(saldo_minutos)


def rotulo_situacao(valor: str | None) -> str:
    return ROTULOS_SITUACAO.get(valor or "", "Indisponível")


def rotulo_status_dia(valor: str | None) -> str:
    if valor in ROTULOS_STATUS_DIA:
        return ROTULOS_STATUS_DIA[valor]
    texto = str(valor or "").replace("_", " ").strip()
    return texto[:1].upper() + texto[1:] if texto else "Indisponível"


def ajustar_larguras(ws) -> None:
    for coluna in ws.columns:
        tamanho = max(len(str(celula.value or "")) for celula in coluna)
        ws.column_dimensions[get_column_letter(coluna[0].column)].width = min(max(tamanho + 2, 12), 42)


def estilizar_cabecalho(ws) -> None:
    preenchimento = PatternFill("solid", fgColor="16845B")
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = preenchimento


def finalizar_planilha(ws) -> None:
    estilizar_cabecalho(ws)
    alinhamento = Alignment(vertical="top", wrap_text=True)
    for linha in ws.iter_rows():
        for celula in linha:
            celula.alignment = alinhamento
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ajustar_larguras(ws)


@router.get("/excel")
def exportar_excel(
    competencia_id: int = Query(...),
    db: Session = Depends(get_db),
) -> StreamingResponse:
    resultado = apurar_competencia(db, competencia_id)
    if not resultado:
        raise HTTPException(status_code=404, detail="Competência não encontrada.")
    db.commit()

    wb = Workbook()
    ws_resumo = wb.active
    ws_resumo.title = "Resumo"
    ws_resumo.append(
        [
            "Empresa",
            "Competência",
            "Funcionário",
            "Atrasos",
            "Extras",
            "Faltas",
            "Atestados",
            "Pendências",
            "Situação",
            "Observações",
            "Horas 100% (feriado)",
        ]
    )
    for item in resultado["resumo"]:
        ws_resumo.append(
            [
                texto_seguro_excel(resultado["empresa"]["nome"]),
                resultado["competencia"]["label"],
                texto_seguro_excel(item["funcionario"]),
                valor_calculado_excel(duracao_excel(item["atrasos_minutos"])),
                valor_calculado_excel(duracao_excel(item["extras_minutos"])),
                valor_calculado_excel(item["faltas"]),
                valor_calculado_excel(item["atestados"]),
                valor_calculado_excel(item["pendencias"]),
                rotulo_situacao(item.get("situacao")),
                texto_seguro_excel(item["observacoes"]),
                valor_calculado_excel(duracao_excel(item.get("horas_feriado_minutos", 0))),
            ]
        )
        ws_resumo.cell(ws_resumo.max_row, 4).number_format = FORMATO_DURACAO_EXCEL
        ws_resumo.cell(ws_resumo.max_row, 5).number_format = FORMATO_DURACAO_EXCEL
        ws_resumo.cell(ws_resumo.max_row, 11).number_format = FORMATO_DURACAO_EXCEL
    finalizar_planilha(ws_resumo)

    ws_marcacoes = wb.create_sheet("Marcações")
    ws_marcacoes.append(
        [
            "Data",
            "Funcionário",
            "Entrada",
            "Saída intervalo",
            "Retorno",
            "Saída",
            "Jornada apurada",
            "Saldo",
            "Status",
            "Conferido",
            "Observações",
            "Horas 100% (feriado)",
        ]
    )
    for item in resultado["marcacoes"]:
        ws_marcacoes.append(
            [
                data_excel(item["data"]),
                texto_seguro_excel(item["funcionario"]),
                hora_excel(item["entrada"]),
                hora_excel(item["saida_almoco"]),
                hora_excel(item["retorno_almoco"]),
                hora_excel(item["saida"]),
                (
                    "Indisponível"
                    if item.get("pendente_calculo")
                    else duracao_excel(item["horas_trabalhadas_minutos"])
                ),
                valor_calculado_excel(saldo_excel(item)),
                texto_seguro_excel(item.get("ocorrencias_rotulo") or rotulo_status_dia(item.get("status_dia"))),
                "Sim" if item["conferido"] else "Não",
                texto_seguro_excel(item["observacoes"]),
                "Indisponível" if item.get("horas_feriado_pendente") else duracao_excel(item.get("horas_feriado_minutos", 0)),
            ]
        )
        linha = ws_marcacoes.max_row
        ws_marcacoes.cell(linha, 1).number_format = FORMATO_DATA_EXCEL
        for coluna in range(3, 7):
            ws_marcacoes.cell(linha, coluna).number_format = FORMATO_HORA_EXCEL
        ws_marcacoes.cell(linha, 7).number_format = FORMATO_DURACAO_EXCEL
        ws_marcacoes.cell(linha, 12).number_format = FORMATO_DURACAO_EXCEL
        celula_saldo = ws_marcacoes.cell(linha, 8)
        celula_saldo.number_format = (
            FORMATO_DURACAO_EXCEL if isinstance(celula_saldo.value, timedelta) else "@"
        )
    finalizar_planilha(ws_marcacoes)

    bancos = [item for item in resultado["resumo"] if item.get("banco_horas")]
    if bancos:
        from app.banco_horas.service import formatar_saldo
        ws_banco = wb.create_sheet("Banco de horas")
        ws_banco.append(["Funcionário", "Saldo anterior", "Créditos da competência", "Débitos da competência",
                         "Saldo final", "Data de desligamento", "Saldo no desligamento", "Consolidação"])
        for item in bancos:
            b = item["banco_horas"]
            ws_banco.append([texto_seguro_excel(item["funcionario"]), formatar_saldo(b["saldo_anterior_minutos"]),
                formatar_saldo(b["creditos_competencia_minutos"]), formatar_saldo(-b["debitos_competencia_minutos"]),
                formatar_saldo(b["saldo_final_minutos"]), b.get("data_demissao"),
                formatar_saldo(b["saldo_demissao_minutos"]) if "saldo_demissao_minutos" in b else None,
                "Fechada" if b["consolidado"] else "Lançamentos atuais; competência ainda aberta"])
        finalizar_planilha(ws_banco)

    arquivo = BytesIO()
    wb.save(arquivo)
    arquivo.seek(0)

    nome_empresa = slug_nome_arquivo(resultado["empresa"]["nome"])
    nome = (
        f"on_ponto_{nome_empresa}_"
        f"{resultado['competencia']['ano']:04d}-{resultado['competencia']['mes']:02d}.xlsx"
    )
    return StreamingResponse(
        arquivo,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome}"'},
    )


@router.get("/impressao", response_class=HTMLResponse)
def relatorio_impressao(
    competencia_id: int = Query(...),
    db: Session = Depends(get_db),
) -> HTMLResponse:
    resultado = apurar_competencia(db, competencia_id)
    if not resultado:
        raise HTTPException(status_code=404, detail="Competência não encontrada.")
    db.commit()

    linhas_resumo = "\n".join(
        f"""
        <tr>
          <td>{escape(item['funcionario'])}</td>
          <td>{escape(str(item['atrasos'] if item['atrasos'] is not None else 'Indisponível'))}</td>
          <td>{escape(str(item['extras'] if item['extras'] is not None else 'Indisponível'))}</td>
          <td>{item['faltas'] if item['faltas'] is not None else 'Indisponível'}</td>
          <td>{item['atestados'] if item['atestados'] is not None else 'Indisponível'}</td>
          <td>{item['pendencias'] if item['pendencias'] is not None else 'Indisponível'}</td>
        </tr>
        """
        for item in resultado["resumo"]
    )

    if resultado["pendencias"]:
        linhas_pendencias = "\n".join(
            f"""
            <tr>
              <td>{escape(item['data'])}</td>
              <td>{escape(item['funcionario'])}</td>
              <td>{escape(item['pendencia_motivo'] or '')}</td>
              <td>{escape(item['observacoes'] or '')}</td>
            </tr>
            """
            for item in resultado["pendencias"]
        )
    else:
        linhas_pendencias = "<tr><td colspan='4'>Sem pendências.</td></tr>"

    linhas_ocorrencias = "".join(
        f"<tr><td>{escape(item['data'])}</td><td>{escape(item['funcionario'])}</td>"
        f"<td>{escape(item.get('ocorrencias_rotulo', ''))}</td>"
        f"<td>{item.get('minutos_abonados', 0)}</td></tr>"
        for item in resultado["marcacoes"] if item.get("ocorrencias")
    )
    secao_ocorrencias = (
        "<section><h2>Ocorrências na apuração</h2><table><thead><tr>"
        "<th>Data</th><th>Funcionário</th><th>Ocorrência</th><th>Minutos abonados</th>"
        "</tr></thead><tbody>" + linhas_ocorrencias + "</tbody></table></section>"
    ) if linhas_ocorrencias else ""

    html = f"""
    <!doctype html>
    <html lang="pt-BR">
    <head>
      <meta charset="utf-8">
      <title>Relatório On Ponto</title>
      <style>
{ESTILO_IMPRESSAO}
      </style>
    </head>
    <body>
      <main>
        <button onclick="window.print()">Imprimir / Salvar PDF</button>
        <header>
          <h1>{escape(resultado['empresa']['nome'])}</h1>
          <div class="meta">
            <span>Competência: {escape(resultado['competencia']['label'])}</span>
            <span>Status: {escape(resultado['competencia']['status'])}</span>
            <span>Gerado em: {escape(resultado['gerado_em'])}</span>
          </div>
        </header>

        <section>
          <h2>Resumo por funcionário</h2>
          <table>
            <thead>
              <tr>
                <th>Funcionário</th>
                <th>Atrasos</th>
                <th>Extras</th>
                <th>Faltas</th>
                <th>Atestados</th>
                <th>Pendências</th>
              </tr>
            </thead>
            <tbody>{linhas_resumo}</tbody>
          </table>
        </section>

        <section>
          <h2>Pendências</h2>
          <table>
            <thead>
              <tr>
                <th>Data</th>
                <th>Funcionário</th>
                <th>Motivo</th>
                <th>Observações</th>
              </tr>
            </thead>
            <tbody>{linhas_pendencias}</tbody>
          </table>
        </section>

        {secao_ocorrencias}
        {''.join('<h2>' + escape(item['funcionario']) + '</h2>' + bloco_banco_horas(item) for item in resultado['resumo'] if item.get('banco_horas'))}
        <footer>On Ponto - relatório imprimível gerado pelo navegador.</footer>
      </main>
    </body>
    </html>
    """
    return HTMLResponse(html)


ESTILO_ESPELHO = ESTILO_IMPRESSAO + """
  .aviso { background:var(--warning-surface); border:1px solid var(--warning); padding:12px; }
  .meta span, td, li { overflow-wrap:anywhere; }
  .espelho-documento { width:210mm; max-width:100%; padding:10mm; box-sizing:border-box; min-height:0; }
  header { margin-bottom:12px; padding-bottom:10px; }
  header h1 { font-size:18px; }
  header h2 { font-size:14px; }
  .meta { font-size:11px; gap:6px 14px; }
  .dias { font-size:9.5px; line-height:1.2; table-layout:fixed; margin:8px 0 12px; }
  .dias th, .dias td { padding:2px; }
  .totais { display:flex; flex-wrap:wrap; gap:6px 16px; margin:12px 0; font-size:11px; }
  .assinaturas { display:flex; gap:24px; margin:32px 0 12px; break-inside:avoid; }
  .assinatura { flex:1; display:flex; flex-wrap:wrap; justify-content:space-between; gap:8px; border-top:1px solid var(--text-secondary); padding-top:6px; font-size:10px; }
  .data-assinatura { white-space:nowrap; }
  .notas { font-size:10px; line-height:1.2; }
  .notas h2 { font-size:12px; }
  .notas ul { padding-left:16px; margin:6px 0; }
  .notas li { margin-bottom:3px; white-space:pre-wrap; }
  .espelho-pagina + .espelho-pagina { margin-top:32px; }
  @page { size:A4 portrait; margin:10mm; }
  @media print {
    main.espelho-documento { width:auto; max-width:none; padding:0; min-height:0; }
    .espelho-pagina + .espelho-pagina { margin-top:0; }
    .espelho-pagina:not(:last-child) { break-after:page; page-break-after:always; }
    thead { display:table-header-group; }
    tr, .totais, .notas li { break-inside:avoid; }
    header { margin-bottom:12px; padding-bottom:10px; }
  }
"""


def bloco_banco_horas(resumo):
    from app.banco_horas.service import formatar_saldo
    banco = resumo.get("banco_horas")
    if not banco:
        return ""
    html = '<section class="totais banco-horas" aria-label="Banco de horas"><span>Saldo do banco de horas ao final da competência: <strong>' + formatar_saldo(banco["saldo_final_minutos"]) + '</strong></span>'
    if not banco["consolidado"]:
        html += '<span>Saldo dos lançamentos consolidados. Esta competência entrará no banco ao fechar.</span>'
    if banco.get("data_demissao"):
        html += '<span><strong>Saldo do banco de horas na data do desligamento (' + escape(banco["data_demissao"]) + '): ' + formatar_saldo(banco["saldo_demissao_minutos"]) + '</strong></span>'
    return html + '</section>'


def renderizar_bloco_espelho(resultado: dict, resumo: dict, dias: list, cargo: str | None) -> str:
    """Mesmo conteúdo individual para emissão avulsa e em lote."""
    def texto(valor):
        return escape(str(valor)) if valor is not None and valor != "" else "—"

    def duracao(minutos):
        return f"{minutos // 60:02d}:{minutos % 60:02d}"

    semana = ("Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom")
    linhas = []
    notas = []
    for item in dias:
        data = date.fromisoformat(item["data"])
        rotulo_data = data.strftime("%d/%m/%Y")
        indisponivel = item.get("pendente_calculo", False)
        valores = [rotulo_data, semana[data.weekday()], item.get("entrada"), item.get("saida_almoco"),
                   item.get("retorno_almoco"), item.get("saida"), rotulo_status_dia(item.get("status_dia")),
                   "Indisponível" if indisponivel else item.get("atraso"),
                   "Indisponível" if indisponivel else item.get("extra"),
                   item.get("ocorrencias_rotulo"), item.get("minutos_abonados") or None]
        celulas = [texto(v) for v in valores]
        if item.get("horas_feriado_pendente"):
            celulas[6] += '<br><small>Horas 100%: Indisponível</small>'
        elif item.get("horas_feriado_minutos", 0):
            celulas[6] += '<br><small>Horas 100%: ' + duracao(item["horas_feriado_minutos"]) + '</small>'
        linhas.append('<tr class="dia">' + "".join(f"<td>{v}</td>" for v in celulas) + "</tr>")
        observacao = item.get("observacoes")
        motivo = item.get("pendencia_motivo") if item.get("pendente") else None
        if motivo or observacao:
            notas.append(f"<li><strong>{rotulo_data}</strong> — "
                         + " · ".join(texto(v) for v in (motivo, observacao) if v) + "</li>")

    pendente = any(item.get("pendente") for item in dias)
    aviso = ('<p class="aviso" role="note">Esta competência ainda possui pendências de conferência '
             '— documento sujeito a revisão.</p>') if pendente and resultado["competencia"]["status"] != "fechada" else ""
    if not dias:
        linhas.append('<tr><td colspan="11">Sem marcações disponíveis para este funcionário nesta competência.</td></tr>')
    calculavel = bool(dias) and not any(item.get("pendente_calculo") or item.get("atraso_minutos") is None
                                      or item.get("extra_minutos") is None for item in dias)
    atrasos = duracao(sum(item["atraso_minutos"] for item in dias)) if calculavel else "Indisponível"
    extras = duracao(sum(item["extra_minutos"] for item in dias)) if calculavel else "Indisponível"
    faltas = sum(item.get("falta", int(item.get("status_dia") == "falta")) for item in dias) if dias else "Indisponível"
    atestados = sum(item.get("atestado", int(item.get("status_dia") == "atestado")) for item in dias) if dias else "Indisponível"
    feriado_pendente = any(item.get("horas_feriado_pendente", False) for item in dias)
    total_feriado = sum(item.get("horas_feriado_minutos", 0) for item in dias)
    feriado_html = ('<span>Horas 100% (feriado): <strong>' + ("Indisponível" if feriado_pendente else duracao(total_feriado)) + '</strong></span>') if total_feriado or feriado_pendente else ""
    com_ocorrencia = sum(bool(item.get("ocorrencias")) for item in dias)
    notas_html = '<section class="notas"><h2>Observações e pendências dos dias</h2><ul>' + "".join(notas) + "</ul></section>" if notas else ""
    return f'''<article class="espelho-pagina" data-funcionario-id="{texto(resumo["funcionario_id"])}">
<header><h1>Espelho de ponto individual</h1><h2>{texto(resultado['empresa']['nome'])}</h2>
<div class="meta"><span>CNPJ: {texto(resultado['empresa'].get('cnpj'))}</span>
<span>Competência: {texto(resultado['competencia']['label'])}</span>
<span>Status: {texto(resultado['competencia']['status'])}</span></div>
<div class="meta"><span>Funcionário: {texto(resumo['funcionario'])}</span>
<span>Código: {texto(resumo.get('codigo'))}</span><span>Cargo (cadastro atual): {texto(cargo)}</span></div></header>
{aviso}
<table class="dias" aria-label="Marcações do funcionário"><colgroup><col style="width:12%"><col style="width:4%"><col span="4" style="width:7%"><col style="width:14%"><col span="2" style="width:7%"><col style="width:20%"><col style="width:8%"></colgroup><thead><tr>
<th>Data</th><th>Dia</th><th>Entrada</th><th>Saída almoço</th><th>Retorno</th><th>Saída</th>
<th>Status do dia</th><th>Atraso</th><th>Extra</th><th>Ocorrência</th><th>Abono (min)</th>
</tr></thead><tbody>{''.join(linhas)}</tbody></table>
<section class="totais" aria-label="Totais do funcionário">
<span>Atrasos: <strong>{atrasos}</strong></span><span>Extras: <strong>{extras}</strong></span>
<span>Faltas: <strong>{faltas}</strong></span><span>Atestados: <strong>{atestados}</strong></span>
<span>Dias com ocorrência: <strong>{com_ocorrencia}</strong></span>{feriado_html}</section>
{bloco_banco_horas(resumo)}
{notas_html}
<section class="assinaturas" aria-label="Assinaturas">
<div class="assinatura">{texto(resumo["funcionario"])}<span class="data-assinatura">Data: ____/____/________</span></div>
<div class="assinatura">{texto(resultado['empresa']['nome'])}<span class="data-assinatura">Data: ____/____/________</span></div>
</section><footer>On Ponto — espelho individual para conferência e assinatura.</footer>
</article>'''


def documento_espelhos(blocos: str, titulo: str) -> HTMLResponse:
    return HTMLResponse(f'''<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8"><title>{escape(titulo)}</title>
<style>{ESTILO_ESPELHO}</style></head><body><main class="espelho-documento">
<button onclick="window.print()">Imprimir / Salvar PDF</button>
{blocos}
</main></body></html>''')


@router.get("/espelho-ponto", response_class=HTMLResponse)
def espelho_ponto(
    competencia_id: int = Query(...),
    funcionario_id: int = Query(...),
    db: Session = Depends(get_db),
) -> HTMLResponse:
    resultado = apurar_competencia(db, competencia_id)
    if resultado is None:
        raise HTTPException(404, "Competência não encontrada.")
    resumo = next((item for item in resultado["resumo"]
                   if item["funcionario_id"] == funcionario_id), None)
    if resumo is None:
        raise HTTPException(404, "Funcionário não encontrado nesta competência.")
    dias = sorted((item for item in resultado["marcacoes"]
                   if item["funcionario_id"] == funcionario_id), key=lambda item: item["data"])
    funcionario = db.get(Funcionario, funcionario_id)
    cargo = funcionario.cargo if funcionario else None
    db.commit()

    bloco = renderizar_bloco_espelho(resultado, resumo, dias, cargo)
    return documento_espelhos(bloco, f"Espelho de ponto — {resumo['funcionario']}")


@router.get("/espelho-ponto-lote", response_class=HTMLResponse)
def espelho_ponto_lote(
    competencia_id: int = Query(...),
    db: Session = Depends(get_db),
) -> HTMLResponse:
    resultado = apurar_competencia(db, competencia_id)
    if resultado is None:
        raise HTTPException(404, "Competência não encontrada.")
    # O resumo preservado determina quem aparece, inclusive quem não tem dias.
    dias_por_funcionario = {}
    for dia in resultado["marcacoes"]:
        dias_por_funcionario.setdefault(dia["funcionario_id"], []).append(dia)
    ids = [item["funcionario_id"] for item in resultado["resumo"]]
    cargos = dict(db.query(Funcionario.id, Funcionario.cargo).filter(Funcionario.id.in_(ids)).all()) if ids else {}
    blocos = [renderizar_bloco_espelho(
        resultado, resumo,
        sorted(dias_por_funcionario.get(resumo["funcionario_id"], []), key=lambda item: item["data"]),
        cargos.get(resumo["funcionario_id"]),
    ) for resumo in resultado["resumo"]]
    db.commit()
    conteudo = "\n".join(blocos) if blocos else '<p role="status">Nenhum funcionário disponível nesta competência.</p>'
    return documento_espelhos(conteudo, f"Espelhos de ponto — {resultado['competencia']['label']}")
