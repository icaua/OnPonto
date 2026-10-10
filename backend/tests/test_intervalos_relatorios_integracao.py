import asyncio
from datetime import date, time, timedelta
from io import BytesIO
from pathlib import Path
import sys
import unittest

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from app.database.session import Base
from app.database.models import Empresa, Escala, Funcionario, Competencia, MarcacaoPonto, LancamentoBancoHoras, EventoCalendario, FechamentoVersao
from app.database.migrations import aplicar_migracoes_compativeis
from app.apuracao.service import apurar_competencia
from app.competencias.routes import fechar_competencia
from app.competencias.schemas import FechamentoCompetenciaRequest
from app.relatorios.routes import exportar_excel, renderizar_bloco_espelho


class IntervalosRelatoriosTest(unittest.TestCase):
    def test_dia_resumo_snapshot_ledger_excel_espelho_usam_os_mesmos_resultados(self):
        engine=create_engine("sqlite://");self.addCleanup(engine.dispose)
        Base.metadata.create_all(engine)
        with Session(engine) as db:
            e=Empresa(nome="FICTÍCIA relatório");db.add(e);db.flush()
            s=Escala(empresa_id=e.id,nome="FICTÍCIA",modo_apuracao="horario_fixo",horario_entrada_prevista=time(9),
                horario_saida_almoco_prevista=time(12),horario_retorno_almoco_prevista=time(13),horario_saida_prevista=time(18),
                tipo_intervalo="fixo",registro_intervalo="pre_assinalado",duracao_intervalo_minutos=60,usa_banco_horas=True,
                politica_horas={"modo":"misto","percentual_folha":50,"percentual_banco":50,"adicional_folha_percentual":60,
                    "ciclo_dias":180,"inicio_ciclo":"2026-09-01"},politica_dsr={"versao":1,"dia_descanso":6,"descontar_feriados":True})
            db.add(s);db.flush()
            f=Funcionario(empresa_id=e.id,nome="FICTÍCIO",escala_id=s.id,data_admissao=date(2026,9,1));db.add(f);db.flush()
            c=Competencia(empresa_id=e.id,mes=9,ano=2026);db.add(c);db.flush()
            for n in range(1,31):
                db.add(MarcacaoPonto(competencia_id=c.id,funcionario_id=f.id,data=date(2026,9,n),
                    status_dia="normal" if n==1 else "falta" if n==2 else "sem_expediente",entrada=time(9) if n==1 else None,
                    saida=time(19) if n==1 else None,batidas_originais='["09:00","19:00"]' if n==1 else None,conferido=True))
            db.add(EventoCalendario(nome="FICTÍCIO",data=date(2026,9,4),tipo="FERIADO",abrangencia="EMPRESA",empresa_id=e.id))
            db.commit();aplicar_migracoes_compativeis(engine)
            aberto=apurar_competencia(db,c.id)
            self.assertEqual((aberto["resumo"][0]["extra_folha_minutos"],aberto["resumo"][0]["extra_banco_minutos"],aberto["resumo"][0]["dsr_desconto_dias"]),(30,30,2))
            self.assertEqual(aberto["marcacoes"][0]["horas_trabalhadas_minutos"],540)
            self.assertIsNone(aberto["marcacoes"][0]["saida_almoco"])
            fechar_competencia(c.id,FechamentoCompetenciaRequest(),db)
            fechado=apurar_competencia(db,c.id)
            self.assertEqual(fechado["resumo"][0]["extra_folha_minutos"],30)
            self.assertEqual(sum(l.minutos for l in db.query(LancamentoBancoHoras) if l.natureza=="credito"),30)
            self.assertEqual(db.query(MarcacaoPonto).filter_by(data=date(2026,9,1)).one().batidas_originais,'["09:00","19:00"]')
            async def conteudo(response): return b"".join([b async for b in response.body_iterator])
            wb=load_workbook(BytesIO(asyncio.run(conteudo(exportar_excel(c.id,db)))),data_only=True)
            try:
                self.assertEqual(wb["Marcações"].cell(2,4).value,None)
                self.assertEqual(wb["Marcações"].cell(2,5).value,None)
                self.assertEqual(wb["Marcações"].cell(2,7).value,timedelta(minutes=540))
                self.assertEqual(wb["Intervalos aplicados"].cell(2,8).value,"Sim")
                self.assertEqual(sum(r[4] for r in list(wb["DSR e feriados"].values)[1:]),2)
                self.assertEqual(wb["Resumo"].cell(1,11).value,"Horas em feriado")
            finally: wb.close()
            html=renderizar_bloco_espelho(fechado,fechado["resumo"][0],fechado["marcacoes"],None)
            self.assertIn("Intervalo previsto (pré-assinalado): 12:00–13:00",html)
            self.assertIn("Desconto DSR/feriados (dias): <strong>2</strong>",html)
            self.assertIn("60%",html)
            self.assertNotIn("Horas 100%",html)
            anterior=db.query(FechamentoVersao).one().snapshot
            db.rollback();aplicar_migracoes_compativeis(engine)
            self.assertEqual(db.query(FechamentoVersao).one().snapshot,anterior)
