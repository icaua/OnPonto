from __future__ import annotations

import sys
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.importadores import txt_id_tempo_maquina, txt_log_relogio, xlsx_cartao_ponto, xlsx_ponto_generico
from app.importadores.deteccao import detectar_adaptadores_txt, detectar_adaptadores_xlsx


FIXTURES = Path(__file__).resolve().parent / "fixtures"


class DeteccaoTxtTests(unittest.TestCase):
    def test_fixture_txt_log_relogio_e_reconhecida_por_exatamente_um_adaptador(self) -> None:
        conteudo = (FIXTURES / "relogio_ficticio.txt").read_bytes()
        encontrados = detectar_adaptadores_txt(conteudo)
        self.assertEqual(encontrados, [txt_log_relogio])

    def test_fixture_id_tempo_maquina_e_reconhecida_por_exatamente_um_adaptador(self) -> None:
        conteudo = (FIXTURES / "relogio_id_tempo_maquina.txt").read_bytes()
        encontrados = detectar_adaptadores_txt(conteudo)
        self.assertEqual(encontrados, [txt_id_tempo_maquina])

    def test_conteudo_sem_relacao_nao_e_reconhecido_por_nenhum_adaptador(self) -> None:
        self.assertEqual(detectar_adaptadores_txt(b"qualquer coisa\nsem cabecalho tabulado"), [])


class DeteccaoXlsxTests(unittest.TestCase):
    def test_fixture_cartao_ponto_e_reconhecida_por_exatamente_um_adaptador(self) -> None:
        workbook = load_workbook(FIXTURES / "cartao_ponto_longo_com_duplicidade.xlsx", data_only=True)
        encontrados = detectar_adaptadores_xlsx(workbook)
        self.assertEqual(encontrados, [xlsx_cartao_ponto])

    def test_layout_largo_e_reconhecido_por_exatamente_um_adaptador(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "NUMERO DE FUNCIONARIO: 007"
        ws["B1"] = "NOME: Fulano de Tal"
        ws["C2"] = "08:00 12:00 13:00 17:00"
        ws["C3"] = 1
        encontrados = detectar_adaptadores_xlsx(wb)
        self.assertEqual(encontrados, [xlsx_ponto_generico])

    def test_workbook_sem_relacao_nao_e_reconhecido_por_nenhum_adaptador(self) -> None:
        wb = Workbook()
        wb.active["A1"] = "Qualquer coisa"
        self.assertEqual(detectar_adaptadores_xlsx(wb), [])


if __name__ == "__main__":
    unittest.main()
