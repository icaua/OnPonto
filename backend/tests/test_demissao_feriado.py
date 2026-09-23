import json
import asyncio
import unittest
from datetime import date, time, timedelta
from io import BytesIO
from openpyxl import load_workbook

import test_admissao_calendario as fixture
from app.apuracao.service import detalhe_marcacao
from app.database.models import Funcionario, MarcacaoPonto
from app.database.migrations import aplicar_migracoes_compativeis, SCHEMA_VERSION
from app.relatorios.routes import exportar_excel, espelho_ponto, espelho_ponto_lote


class DemissaoFeriadoTest(unittest.TestCase):
    setUp = fixture.AdmissaoCalendarioTest.setUp
    tearDown = fixture.AdmissaoCalendarioTest.tearDown
    apurar = fixture.AdmissaoCalendarioTest.apurar
    evento = fixture.AdmissaoCalendarioTest.evento
    batida = fixture.AdmissaoCalendarioTest.batida
    dia = fixture.AdmissaoCalendarioTest.dia

    def excel(self):
        response = exportar_excel(self.c.id,self.db)
        async def read():
            return b''.join([chunk async for chunk in response.body_iterator])
        workbook = load_workbook(BytesIO(asyncio.run(read())))
        self.addCleanup(workbook.close)
        return workbook

    def test_feriado_sem_batidas_nao_gera_pendencia_inclusive_manual_sem_escala(self):
        self.f.escala = None
        m=self.batida(entrada=None,batidas_originais=None,status_dia='feriado',conferido=False)
        d=detalhe_marcacao(self.f,m)
        for campo in ['horas_trabalhadas_minutos','horas_feriado_minutos','jornada_prevista_minutos','atraso_minutos','extra_minutos']:
            self.assertEqual(d[campo],0,campo)
        self.assertFalse(d['pendente'])
        self.assertFalse(m.conferido)

    def test_feriado_batidas_completas_duas_ou_quatro_menor_maior_jornada(self):
        m=self.batida(status_dia='feriado',entrada=time(8),saida=time(12),conferido=True)
        for saida,almoco,retorno,esperado in [(time(12),None,None,240),(time(20),None,None,720),(time(18),time(12),time(13),540)]:
            with self.subTest(esperado=esperado):
                m.saida=saida;m.saida_almoco=almoco;m.retorno_almoco=retorno
                d=detalhe_marcacao(self.f,m)
                self.assertEqual(d['horas_trabalhadas_minutos'],esperado)
                self.assertEqual(d['horas_feriado_minutos'],esperado)
                self.assertEqual(d['horas_feriado'],f'{esperado//60:02}:{esperado%60:02}')
                self.assertEqual(d['atraso_minutos'],0)
                self.assertEqual(d['extra_minutos'],0)
                self.assertFalse(d['pendente'])

    def test_feriado_parcial_e_sequencia_invalida_nao_inventam_horas(self):
        m=self.batida(status_dia='feriado',conferido=True)
        for entrada,saida,almoco,retorno in [(time(8),None,None,None),(time(8),time(17),time(12),None),(time(17),time(8),None,None)]:
            m.entrada=entrada;m.saida=saida;m.saida_almoco=almoco;m.retorno_almoco=retorno
            d=detalhe_marcacao(self.f,m)
            self.assertTrue(d['pendente_calculo'])
            self.assertTrue(d['horas_feriado_pendente'])
            self.assertEqual(d['pendencia_tipo'],'batidas_insuficientes')
            self.assertIn('feriado',d['pendencia_motivo'])
            self.assertEqual(d['horas_feriado_minutos'],0)
            self.assertEqual(d['horas_trabalhadas_minutos'],0)

    def test_feriado_com_originais_nao_interpretados_requer_revisao(self):
        m=self.batida(status_dia='feriado',entrada=None)
        d=detalhe_marcacao(self.f,m)
        self.assertTrue(d['pendente'])
        self.assertEqual(d['horas_feriado_minutos'],0)

    def test_resumo_soma_feriados_separado_por_pessoa_sem_tolerancia(self):
        self.evento(data=date(2026,6,15))
        self.evento(data=date(2026,6,16))
        self.batida(entrada=time(8),saida=time(12))
        self.batida(data=date(2026,6,16),entrada=time(8),saida=time(18))
        outra=Funcionario(empresa_id=self.empresa.id,nome='Segunda',codigo='2',escala_id=self.escala.id)
        self.db.add(outra);self.db.flush()
        self.batida(funcionario_id=outra.id,entrada=time(8),saida=time(8,3))
        r=self.apurar()
        resumo={x['funcionario_id']:x for x in r['resumo']}
        self.assertEqual(resumo[self.f.id]['horas_feriado_minutos'],840)
        self.assertEqual(resumo[self.f.id]['horas_feriado'],'14:00')
        self.assertEqual(resumo[outra.id]['horas_feriado_minutos'],3)
        self.assertTrue(all(x['horas_feriado_minutos']==0 for x in r['marcacoes'] if x['status_dia']!='feriado'))

    def test_feriado_fora_vinculo_nao_gera_horas_100(self):
        self.f.data_demissao=date(2026,6,14)
        self.evento(data=date(2026,6,15))
        self.batida(saida=time(19))
        d=self.dia(self.apurar(),15)
        self.assertEqual(d['horas_feriado_minutos'],0)
        self.assertEqual(d['pendencia_tipo'],'marcacao_posterior_demissao')

    def test_excel_numerico_e_espelhos_individual_lote_mostram_100_separado(self):
        self.batida(status_dia='feriado',entrada=time(8),saida=time(12),conferido=True)
        wb=self.excel()
        resumo,marcacoes=wb['Resumo'],wb['Marcações']
        self.assertEqual(resumo['K1'].value,'Horas 100% (feriado)')
        self.assertEqual(resumo['K2'].value,timedelta(hours=4))
        self.assertEqual(resumo['K2'].number_format,'[h]:mm')
        linha=next(row for row in marcacoes.iter_rows(min_row=2) if row[0].value.day==15)
        self.assertEqual(linha[11].value,timedelta(hours=4))
        self.assertEqual(linha[11].number_format,'[h]:mm')
        for resposta in [espelho_ponto(self.c.id,self.f.id,self.db),espelho_ponto_lote(self.c.id,self.db)]:
            html=resposta.body.decode()
            self.assertIn('Horas 100%: 04:00',html)
            self.assertIn('Horas 100% (feriado): <strong>04:00</strong>',html)
            self.assertNotIn('<th>Horas 100%',html)

    def test_feriado_incompleto_total_indisponivel_e_relatorio_avisa(self):
        self.batida(status_dia='feriado')
        r=self.apurar()
        self.assertIsNone(r['resumo'][0]['horas_feriado_minutos'])
        self.assertEqual(self.excel()['Resumo']['K2'].value,'Indisponível')
        html=espelho_ponto(self.c.id,self.f.id,self.db).body.decode()
        self.assertIn('Horas 100%: Indisponível',html)
        self.assertIn('Horas 100% (feriado): <strong>Indisponível</strong>',html)

    def test_snapshot_antigo_sem_novos_campos_leitura_zero_sem_regravar(self):
        self.batida(status_dia='feriado',entrada=time(8),saida=time(12),conferido=True)
        snapshot=self.apurar()
        for item in snapshot['resumo']+snapshot['marcacoes']:
            for campo in ['horas_feriado_minutos','horas_feriado','horas_feriado_pendente']:
                item.pop(campo,None)
        snapshot['competencia']['status']='fechada'
        self.c.status='fechada';self.c.apuracao_fechada=json.dumps(snapshot)
        original=self.c.apuracao_fechada
        self.assertEqual(self.apurar(),snapshot)
        self.assertEqual(self.excel()['Resumo']['K2'].value,timedelta(0))
        for resposta in [espelho_ponto(self.c.id,self.f.id,self.db),espelho_ponto_lote(self.c.id,self.db)]:
            self.assertNotIn('Horas 100%',resposta.body.decode())
        self.assertEqual(self.c.apuracao_fechada,original)

    def test_demissao_15_junho_inclusiva(self):
        self.f.data_demissao = date(2026,6,15)
        r = self.apurar()
        self.assertEqual(r['resumo'][0]['dias_processados'],15)
        self.assertEqual(len(r['marcacoes']),15)
        self.assertTrue(all(m['data'] <= '2026-06-15' for m in r['pendencias']))

    def test_admissao_e_demissao_no_mes_com_limites_inclusivos(self):
        self.f.data_admissao = date(2026,6,10)
        self.f.data_demissao = date(2026,6,15)
        r = self.apurar()
        self.assertEqual([m['data'] for m in r['marcacoes']],[f'2026-06-{d:02}' for d in range(10,16)])
        self.assertEqual(r['resumo'][0]['dias_processados'],6)

    def test_batida_posterior_preservada_e_avisada_sem_jornada(self):
        self.f.data_demissao = date(2026,6,14)
        m = self.batida(saida=time(19),conferido=True)
        r = self.apurar()
        d = self.dia(r,15)
        self.assertEqual(d['pendencia_motivo'],'Marcação posterior à data de demissão')
        self.assertEqual(d['pendencia_tipo'],'marcacao_posterior_demissao')
        self.assertEqual(d['status_dia'],'fora_vinculo')
        for campo in ['horas_trabalhadas_minutos','jornada_prevista_minutos','atraso_minutos','extra_minutos','falta']:
            self.assertEqual(d[campo],0,campo)
        self.assertEqual(m.batidas_originais,'["08:02"]')
        self.assertEqual(r['resumo'][0]['dias_processados'],14)
        self.assertEqual(r['resumo'][0]['observacoes'],'Existem marcações posteriores à demissão')

    def test_demissao_antes_competencia_sem_pendencias_artificiais(self):
        self.f.data_demissao = date(2026,5,31)
        r = self.apurar()
        self.assertEqual(r['marcacoes'],[])
        self.assertEqual(r['resumo'][0]['dias_processados'],0)
        self.assertEqual(r['resumo'][0]['pendencias'],0)
        self.assertEqual(r['resumo'][0]['situacao'],'fora_vinculo')

    def test_editar_demissao_limpa_so_placeholders_e_limpar_restaura_calendario(self):
        self.apurar()
        self.f.data_demissao = date(2026,6,15)
        self.assertEqual(len(self.apurar()['marcacoes']),15)
        self.assertEqual(self.db.query(MarcacaoPonto).count(),15)
        self.f.data_demissao = None
        self.assertEqual(len(self.apurar()['marcacoes']),30)

    def test_demissao_nao_altera_snapshot_fechado(self):
        snapshot = self.apurar()
        snapshot['competencia']['status']='fechada'
        self.c.status='fechada'
        self.c.apuracao_fechada=json.dumps(snapshot)
        texto = self.c.apuracao_fechada
        self.f.data_demissao=date(2026,6,15)
        self.assertEqual(self.apurar(),snapshot)
        self.assertEqual(self.c.apuracao_fechada,texto)

    def test_migracao_7_para_8_nullable_preserva_admissao(self):
        self.f.data_admissao=date(2026,6,10)
        self.db.commit()
        self.db.close()
        with self.engine.begin() as c:
            c.exec_driver_sql('ALTER TABLE funcionarios DROP COLUMN data_demissao')
            c.exec_driver_sql('PRAGMA user_version = 7')
        aplicar_migracoes_compativeis(self.engine)
        aplicar_migracoes_compativeis(self.engine)
        with self.engine.connect() as c:
            self.assertEqual(c.exec_driver_sql('SELECT data_admissao,data_demissao FROM funcionarios').one(),('2026-06-10',None))
            self.assertEqual(c.exec_driver_sql('PRAGMA user_version').scalar_one(),SCHEMA_VERSION)
            self.assertEqual(c.exec_driver_sql('PRAGMA foreign_key_check').all(),[])
