import io
import json
import unittest
from datetime import date, timedelta

from fastapi import HTTPException
from pydantic import ValidationError
from openpyxl import load_workbook

from test_banco_horas import BancoFixture
from app.apuracao.service import apurar_competencia
from app.banco_horas.regras import PoliticaHoras, ciclo
from app.banco_horas.service import extrato, reconciliar, tratamentos_pendentes
from app.database.models import LancamentoBancoHoras, CompensacaoBancoHoras
from app.database.migrations import aplicar_migracoes_compativeis
from app.escalas.routes import atualizar
from app.escalas.schemas import EscalaUpdate


def politica(modo="misto", **kw):
    folha = {"folha": 100, "banco": 0, "misto": 50}[modo]
    return dict(modo=modo, percentual_folha=folha, percentual_banco=100-folha,
        adicional_folha_percentual=50, ciclo_dias=180, inicio_ciclo="2026-09-01", **kw)


class PoliticasHorasTest(BancoFixture, unittest.TestCase):
    def configurar(self, modo="misto", **kw):
        self.escala.politica_horas = politica(modo, **kw)
        self.escala.usa_banco_horas = modo != "folha"
        self.db.commit()
        self.dias()

    def apurar(self):
        return apurar_competencia(self.db, self.competencia.id)

    def test_destinos_e_atraso_separado(self):
        for modo, esperado in [("folha", (120, 0, 0)), ("banco", (0, 120, 60)), ("misto", (60, 60, 60))]:
            with self.subTest(modo=modo):
                self.configurar(modo)
                self.batida(extra=120, atraso=60)
                item = self.apurar()["resumo"][0]
                self.assertEqual(item["extra_apurada_minutos"], 120)
                self.assertEqual(tuple(item[c] for c in ("extra_folha_minutos", "extra_banco_minutos", "debito_banco_minutos")), esperado)
                self.assertEqual(item["atrasos_minutos"], 60)
                self.assertEqual(item["extra_folha_minutos"] + item["extra_banco_base_minutos"], 120)
                self.assertEqual(item["adicionais_folha"], [] if modo == "banco" else [{"percentual": "50", "minutos": esperado[0]}])

    def test_ledger_preserva_credito_e_debito_mesmo_saldo_zero(self):
        self.configurar(); self.batida(extra=120, atraso=60); self.fechar()
        itens = self.db.query(LancamentoBancoHoras).all()
        self.assertEqual([(i.natureza, i.minutos) for i in itens], [("credito", 60), ("debito", 60)])
        self.assertEqual(self.db.query(CompensacaoBancoHoras).one().minutos, 60)
        self.assertEqual(itens[0].ciclo_inicio, date(2026,9,1))
        self.assertEqual(itens[0].ciclo_fim, date(2027,2,27))
        self.assertEqual(itens[0].data_vencimento, itens[0].ciclo_fim)
        self.assertEqual(json.loads(itens[0].politica_aplicada_json)["politica"]["percentual_banco"], "50")

    def test_snapshot_regra_totais_saldo_nao_mudam_com_futuro(self):
        self.configurar(); self.batida(extra=120); self.fechar()
        antes = self.apurar()
        atualizar(self.escala.id, EscalaUpdate(politica_horas=politica("banco")), self.db)
        self.lancar(natureza="credito", minutos=77); self.db.commit()
        depois = self.apurar()
        self.assertEqual(antes, depois)
        regra = depois["resumo"][0]["politicas_horas_aplicadas"][0]["politica"]
        self.assertEqual((regra["percentual_folha"], regra["adicional_folha_percentual"]), ("50", "50"))
        self.assertEqual(depois["resumo"][0]["banco_horas"]["saldo_final_minutos"], 60)
        self.assertEqual(extrato(self.db, self.funcionario.id)["saldo_minutos"], 137)

    def test_domingo_e_feriado_fator_parcela_ou_integral_explicito(self):
        for base, esperado in [("parcela_banco", 90), ("extra_integral", 180)]:
            for tipo, dia in [("domingo", 6), ("feriado", 7), ("sabado", 5), ("normal", 1)]:
                with self.subTest(base=base, tipo=tipo):
                    self.configurar(**{tipo: {"fator_banco": "1.5", "base_fator": base}})
                    m = self.batida(dia=dia, extra=120)
                    if tipo != "normal":
                        from datetime import time
                        m.entrada=time(8); m.saida=time(10); m.saida_almoco=None; m.retorno_almoco=None
                    if tipo == "feriado":
                        m.status_dia="feriado"
                    self.db.commit()
                    item = next(d for d in self.apurar()["marcacoes"] if d["id"] == m.id)
                    self.assertEqual(item["extra_apurada_minutos"], 120)
                    self.assertEqual(item["extra_banco_minutos"], esperado)
                    self.assertEqual(item["extra_folha_minutos"], 60)

    def test_ambiguidade_bloqueia_fechamento_sem_inventar_zero(self):
        self.configurar(normal={"fator_banco": "1.5"}); self.batida(extra=120)
        apuracao = self.apurar()
        item = apuracao["marcacoes"][0]
        self.assertEqual(item["pendencia_tipo"], "politica_requer_decisao")
        self.assertEqual(item["extra_apurada_minutos"], 120)
        self.assertIsNone(item["extra_banco_minutos"])
        self.assertFalse(apuracao["fechamento"]["pode_fechar"])
        with self.assertRaises(HTTPException): self.fechar()
        self.assertEqual(self.db.query(LancamentoBancoHoras).count(), 0)

    def test_arredondamento_acumulado_conserva_e_e_deterministico(self):
        self.configurar(normal={"fator_banco": "1.5", "base_fator": "parcela_banco"})
        for dia in (1, 2, 3, 4): self.batida(dia=dia, extra=1)
        a = self.apurar(); b = self.apurar()
        a.pop("gerado_em"); b.pop("gerado_em")
        self.assertEqual(a, b)
        item = a["resumo"][0]
        self.assertEqual(tuple(item[c] for c in ("extra_apurada_minutos", "extra_folha_minutos", "extra_banco_base_minutos", "extra_banco_minutos")), (4,2,2,3))
        for d in a["marcacoes"]:
            self.assertEqual(d["extra_apurada_minutos"], d["extra_folha_minutos"] + d["extra_banco_base_minutos"])

    def test_validacao_percentuais_modo_e_ciclo(self):
        for mudanca in [{"percentual_folha": 49}, {"modo": "folha"}, {"inicio_ciclo": None}, {"ciclo_dias": 0}, {"normal": {"fator_banco": -1}}]:
            with self.subTest(mudanca=mudanca), self.assertRaises(ValidationError):
                PoliticaHoras.model_validate(politica() | mudanca)
        self.assertEqual(ciclo(PoliticaHoras.model_validate(politica()), date(2027,2,28)), (date(2027,2,28), date(2027,8,26)))

    def test_politica_independe_do_prazo_legado_e_booleano_nao_apaga(self):
        self.empresa.prazo_compensacao_banco_horas_dias=None; self.db.commit()
        atualizar(self.escala.id, EscalaUpdate(politica_horas=politica()), self.db)
        self.assertTrue(self.escala.usa_banco_horas)
        atualizar(self.escala.id, EscalaUpdate(nome="Nome novo"), self.db)
        self.assertEqual(self.escala.politica_horas["modo"], "misto")
        with self.assertRaises(HTTPException):
            atualizar(self.escala.id, EscalaUpdate(usa_banco_horas=False), self.db)

    def test_debito_negativo_permitido_ou_bloqueado(self):
        self.configurar(permite_saldo_negativo=False); self.batida(extra=0, atraso=60)
        self.assertFalse(self.apurar()["fechamento"]["pode_fechar"])
        with self.assertRaises(HTTPException): self.fechar()
        self.assertEqual(self.db.query(LancamentoBancoHoras).count(), 0)
        self.escala.politica_horas=politica(permite_saldo_negativo=True); self.db.commit()
        self.fechar()
        self.assertEqual(extrato(self.db, self.funcionario.id)["saldo_minutos"], -60)

    def test_saldos_vencidos_sinalizados_sem_consumo_por_outro_ciclo(self):
        self.configurar(fim_ciclo_credor="pagar", desligamento_credor="pagar", desligamento_devedor="decidir")
        self.batida(extra=120); self.fechar()
        pendentes=tratamentos_pendentes(self.db, self.funcionario, date(2027,2,28))
        self.assertEqual((pendentes[0]["tratamento"], pendentes[0]["minutos"]), ("pagar",60))
        debito=self.lancar(natureza="debito", minutos=60, dia=date(2027,3,1),
            ciclo_inicio=date(2027,2,28), ciclo_fim=date(2027,8,26))
        reconciliar(self.db, self.funcionario.id)
        self.assertEqual(self.db.query(CompensacaoBancoHoras).count(), 0)
        self.funcionario.data_demissao=date(2026,9,30); self.db.commit()
        self.assertEqual(tratamentos_pendentes(self.db,self.funcionario,date(2026,9,30))[0]["evento"], "desligamento")

    def test_migracao_11_aditiva_idempotente_preserva_dados(self):
        self.dias(); self.fechar(); antes=self.competencia.apuracao_fechada
        escala_id, competencia_id = self.escala.id, self.competencia.id
        self.db.close()
        with self.engine.begin() as c:
            # SQLite não remove isoladamente uma coluna que integra uma FK de tabela.
            for coluna in ("politica_aplicada_json", "ciclo_inicio", "ciclo_fim"):
                c.exec_driver_sql("ALTER TABLE lancamentos_banco_horas DROP COLUMN " + coluna)
            c.exec_driver_sql("ALTER TABLE escalas DROP COLUMN politica_horas_json")
            c.exec_driver_sql("PRAGMA user_version=10")
        aplicar_migracoes_compativeis(self.engine); aplicar_migracoes_compativeis(self.engine)
        from app.database.models import Escala, Competencia
        self.assertIsNone(self.db.get(Escala, escala_id).politica_horas)
        self.assertEqual(self.db.get(Competencia, competencia_id).apuracao_fechada, antes)

    def test_relatorios_separam_componentes_e_politica(self):
        import asyncio
        from app.relatorios.routes import exportar_excel, espelho_ponto, relatorio_impressao
        self.configurar(); self.batida(extra=120, atraso=60); self.fechar()
        resposta = exportar_excel(self.competencia.id, self.db)
        async def consumir():
            return b"".join([parte async for parte in resposta.body_iterator])
        wb = load_workbook(io.BytesIO(asyncio.run(consumir())))
        self.addCleanup(wb.close)
        self.assertIn("Políticas aplicadas", wb.sheetnames)
        linha=list(wb["Distribuição de HE"].values)[1]
        self.assertEqual(linha[1:6], (timedelta(hours=2), timedelta(hours=1), timedelta(hours=1), timedelta(hours=1), timedelta(hours=1)))
        self.assertEqual(linha[6], "01:00 a 50%")
        for resposta in (espelho_ponto(self.competencia.id,self.funcionario.id,self.db), relatorio_impressao(self.competencia.id,self.db)):
            html=resposta.body.decode()
            self.assertIn("Distribuição das horas extras", html)
            self.assertIn("01:00 a 50%", html)
            self.assertIn("Saldo inicial", html)

    def test_leitura_politica_nao_persiste_nem_gera_ledger(self):
        self.configurar(); self.batida(extra=120)
        from sqlalchemy import text
        antes=list(self.db.execute(text("SELECT * FROM competencias")))
        self.apurar(); self.apurar()
        self.assertEqual(antes,list(self.db.execute(text("SELECT * FROM competencias"))))
        self.assertFalse(self.db.new); self.assertFalse(self.db.dirty)
        self.assertEqual(self.db.query(LancamentoBancoHoras).count(),0)

    def test_folga_negativa_e_ajuste_respeitam_politica_do_ciclo(self):
        from app.ocorrencias.routes import criar
        from app.ocorrencias.schemas import OcorrenciaCreate
        from app.banco_horas.routes import ajustar
        from app.banco_horas.schemas import AjusteCreate
        self.configurar(permite_saldo_negativo=True)
        criar(OcorrenciaCreate(funcionario_id=self.funcionario.id,tipo="FOLGA_COMPENSATORIA",
            data_inicio=date(2026,9,1),data_fim=date(2026,9,1)),self.db)
        self.fechar()
        self.assertEqual(extrato(self.db,self.funcionario.id)["saldo_minutos"],-480)
        resposta=ajustar(AjusteCreate(funcionario_id=self.funcionario.id,natureza="credito",minutos=480,
            data_referencia=date(2026,9,2),observacao="Ajuste fictício documentado"),self.db)
        item=self.db.get(LancamentoBancoHoras,resposta["id"])
        self.assertEqual(item.ciclo_inicio,date(2026,9,1))
        self.assertEqual(extrato(self.db,self.funcionario.id)["saldo_minutos"],0)

    def test_ajuste_referencia_preserva_ciclo_antigo_apos_mudar_politica(self):
        from app.banco_horas.routes import ajustar
        from app.banco_horas.schemas import AjusteCreate
        self.configurar(); self.batida(extra=120); self.fechar()
        original=self.db.query(LancamentoBancoHoras).one()
        dados=politica(); dados.update(ciclo_dias=90,inicio_ciclo="2026-08-01")
        atualizar(self.escala.id,EscalaUpdate(politica_horas=dados),self.db)
        resposta=ajustar(AjusteCreate(funcionario_id=self.funcionario.id,natureza="debito",minutos=60,
            data_referencia=date(2026,9,2),lancamento_referencia_id=original.id,observacao="Quitação fictícia do saldo original"),self.db)
        novo=self.db.get(LancamentoBancoHoras,resposta["id"])
        self.assertEqual(novo.ciclo_inicio,original.ciclo_inicio)
        self.assertEqual(novo.politica_aplicada_json,original.politica_aplicada_json)
        self.assertEqual(self.db.query(CompensacaoBancoHoras).one().minutos,60)
        self.assertEqual(extrato(self.db,self.funcionario.id)["saldo_minutos"],0)
        with self.assertRaises(HTTPException):
            ajustar(AjusteCreate(funcionario_id=self.funcionario.id,natureza="debito",minutos=1,
                data_referencia=date(2026,8,1),lancamento_referencia_id=original.id,observacao="Data fora do ciclo"),self.db)

    def test_saldos_opostos_em_ciclos_distintos_nao_autorizam_desativacao(self):
        self.configurar(); self.batida(extra=120); self.fechar()
        self.lancar(natureza="debito",minutos=60,dia=date(2026,8,1),ciclo_inicio=date(2026,3,5),ciclo_fim=date(2026,8,31))
        reconciliar(self.db,self.funcionario.id); self.db.commit()
        self.assertEqual(extrato(self.db,self.funcionario.id)["saldo_minutos"],0)
        with self.assertRaises(HTTPException):
            atualizar(self.escala.id,EscalaUpdate(politica_horas=politica("folha")),self.db)

    def test_legado_sem_banco_nao_muda_calculo_ou_jornada(self):
        self.escala.usa_banco_horas=False; self.db.commit()
        self.dias(); self.batida(extra=120,atraso=60)
        item=self.apurar()["resumo"][0]
        self.assertEqual((item["extras_minutos"],item["atrasos_minutos"]),(60,60))
        self.assertNotIn("banco_horas",item)
        self.assertEqual(item["politicas_horas_aplicadas"],[])

    def test_adicional_ausente_bloqueia_e_totais_parciais_nao_parecem_completos(self):
        self.configurar()
        self.batida(dia=1,extra=120)
        self.batida(dia=2,extra=60)
        p=politica(); p["adicional_folha_percentual"]=None
        self.escala.politica_horas=p;self.db.commit()
        resultado=self.apurar()
        self.assertFalse(resultado["fechamento"]["pode_fechar"])
        self.assertIsNone(resultado["resumo"][0]["adicionais_folha"])
        self.assertEqual(resultado["resumo"][0]["extra_apurada_minutos"],180)


if __name__ == "__main__": unittest.main()
