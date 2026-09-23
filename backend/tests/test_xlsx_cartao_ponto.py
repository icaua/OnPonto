from __future__ import annotations

import sys
import unittest
from pathlib import Path

from openpyxl import Workbook


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.importadores import xlsx_cartao_ponto as adaptador
from app.importadores.xlsx_cartao_ponto import (
    ErroImportacaoXlsxCartaoPonto,
    detectar,
    parse_xlsx_cartao_ponto,
)
from openpyxl import load_workbook


FIXTURE = Path(__file__).resolve().parent / "fixtures" / "cartao_ponto_longo_com_duplicidade.xlsx"


def _workbook_layout_generico() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "NUMERO DE FUNCIONÁRIO: 7"
    ws["A2"] = "NOME: Fulano de Tal"
    ws["A3"] = "1"
    ws["A4"] = "08:00"
    return wb


class DetectarTests(unittest.TestCase):
    def test_reconhece_a_fixture_real(self) -> None:
        workbook = load_workbook(FIXTURE, data_only=True)
        self.assertTrue(detectar(workbook))

    def test_nao_reconhece_o_layout_largo_do_outro_adaptador_xlsx(self) -> None:
        self.assertFalse(detectar(_workbook_layout_generico()))

    def test_nome_do_adaptador(self) -> None:
        self.assertEqual(adaptador.NOME_ADAPTER, "xlsx_cartao_ponto")


class ParseFixtureRealTests(unittest.TestCase):
    def test_arquivo_real_e_interpretado_por_completo(self) -> None:
        resultado = parse_xlsx_cartao_ponto(
            str(FIXTURE), arquivo_nome="cartao.xlsx", mes=7, ano=2026, funcionarios=()
        )
        self.assertEqual(resultado["total_linhas_validas"], 80)
        self.assertEqual(len(resultado["registros"]), 80)

    def test_dia_com_seis_batidas_e_duplicidade_gera_pendencia_sem_inventar_interpretacao(self) -> None:
        resultado = parse_xlsx_cartao_ponto(
            str(FIXTURE), arquivo_nome="cartao.xlsx", mes=7, ano=2026, funcionarios=()
        )
        registro = next(
            item
            for item in resultado["registros"]
            if item["funcionario"]["codigo_origem"] == "6009654958" and item["data"] == "2026-07-24"
        )
        self.assertEqual(
            registro["batidas_originais"],
            ["08:02", "08:02", "11:59", "12:00", "12:59", "16:58"],
        )
        self.assertTrue(all(valor is None for valor in registro["interpretacao"].values()))
        self.assertEqual(registro["status"], "conferir")
        self.assertIn("Mais de quatro marcações encontradas", registro["pendencias"])

    def test_dia_normal_com_quatro_batidas_preenche_interpretacao(self) -> None:
        resultado = parse_xlsx_cartao_ponto(
            str(FIXTURE), arquivo_nome="cartao.xlsx", mes=7, ano=2026, funcionarios=()
        )
        registro = next(
            item
            for item in resultado["registros"]
            if item["funcionario"]["codigo_origem"] == "6009654958" and item["data"] == "2026-07-23"
        )
        self.assertEqual(
            registro["interpretacao"],
            {"entrada": "08:03", "saida_intervalo": "12:00", "retorno_intervalo": "12:58", "saida": "15:57"},
        )
        self.assertEqual(registro["status"], "nao_conferido")

    def test_sobrenome_com_hifen_usa_apenas_o_nome(self) -> None:
        resultado = parse_xlsx_cartao_ponto(
            str(FIXTURE), arquivo_nome="cartao.xlsx", mes=7, ano=2026, funcionarios=()
        )
        registro = resultado["registros"][0]
        self.assertNotIn("-", registro["funcionario"]["nome_origem"])

    def test_funcionario_cadastrado_por_codigo_e_selecionado(self) -> None:
        resultado = parse_xlsx_cartao_ponto(
            str(FIXTURE),
            arquivo_nome="cartao.xlsx",
            mes=7,
            ano=2026,
            funcionarios=({"id": 99, "codigo": "6009654958", "nome": "Gaby"},),
        )
        registro = next(
            item
            for item in resultado["registros"]
            if item["funcionario"]["codigo_origem"] == "6009654958" and item["data"] == "2026-07-23"
        )
        self.assertTrue(registro["funcionario"]["encontrado"])
        self.assertEqual(registro["funcionario"]["id"], 99)
        self.assertTrue(registro["selecionado"])

    def test_data_fora_da_competencia_e_desmarcada(self) -> None:
        resultado = parse_xlsx_cartao_ponto(
            str(FIXTURE), arquivo_nome="cartao.xlsx", mes=8, ano=2026, funcionarios=()
        )
        self.assertTrue(all(item["fora_da_competencia"] for item in resultado["registros"]))
        self.assertTrue(all(not item["selecionado"] for item in resultado["registros"]))


class ParseErroTests(unittest.TestCase):
    def test_workbook_sem_cabecalho_esperado_gera_erro_legivel(self) -> None:
        wb = _workbook_layout_generico()
        caminho = Path(self._tmp_path())
        wb.save(caminho)
        with self.assertRaises(ErroImportacaoXlsxCartaoPonto) as contexto:
            parse_xlsx_cartao_ponto(str(caminho), arquivo_nome="teste.xlsx", mes=7, ano=2026, funcionarios=())
        self.assertEqual(contexto.exception.codigo, "cabecalho_incompativel")
        caminho.unlink(missing_ok=True)

    def _tmp_path(self) -> str:
        import tempfile

        fd, caminho = tempfile.mkstemp(suffix=".xlsx")
        import os

        os.close(fd)
        return caminho


if __name__ == "__main__":
    unittest.main()
