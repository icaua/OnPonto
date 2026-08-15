from __future__ import annotations

import json
import os
import socket
import subprocess
import sys
import tempfile
import time
import unittest
from datetime import date, datetime, time as horario, timedelta
from io import BytesIO
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook


BACKEND_DIR = Path(__file__).resolve().parents[1]

CABECALHO_RESUMO = [
    "Empresa",
    "Competência",
    "Funcionário",
    "Atrasos",
    "Extras",
    "Faltas",
    "Atestados",
    "Pendências",
    "Situação",
    "Observações",
]

CABECALHO_MARCACOES = [
    "Data",
    "Funcionário",
    "Entrada",
    "Saída intervalo",
    "Retorno",
    "Saída",
    "Jornada apurada",
    "Saldo",
    "Status",
    "Conferido",
    "Observações",
]


class RelatorioExcelApiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temp_dir = tempfile.TemporaryDirectory(prefix="onponto-relatorio-excel-")
        cls.addClassCleanup(cls.temp_dir.cleanup)

        raiz = Path(cls.temp_dir.name)
        banco = raiz / "relatorio-excel.test.db"
        uploads = raiz / "uploads"
        uploads.mkdir()

        with socket.socket() as servidor:
            servidor.bind(("127.0.0.1", 0))
            cls.porta = servidor.getsockname()[1]

        ambiente = os.environ.copy()
        ambiente["ONPONTO_DATABASE_URL"] = f"sqlite:///{banco}"
        ambiente["ONPONTO_UPLOADS_DIR"] = str(uploads)
        cls.processo = subprocess.Popen(
            [
                sys.executable,
                "-m",
                "uvicorn",
                "app.main:app",
                "--host",
                "127.0.0.1",
                "--port",
                str(cls.porta),
                "--log-level",
                "warning",
            ],
            cwd=BACKEND_DIR,
            env=ambiente,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
        )
        cls.addClassCleanup(cls.encerrar_api)
        cls.base_url = f"http://127.0.0.1:{cls.porta}"

        limite = time.monotonic() + 10
        while time.monotonic() < limite:
            if cls.processo.poll() is not None:
                saida = cls.processo.stdout.read() if cls.processo.stdout else ""
                raise RuntimeError(f"A API de teste encerrou antes de iniciar:\n{saida}")
            try:
                with urlopen(cls.base_url + "/", timeout=0.5) as resposta:
                    if resposta.status == 200:
                        return
            except (URLError, TimeoutError):
                time.sleep(0.1)
        raise RuntimeError("A API de teste não iniciou dentro do tempo esperado.")

    @classmethod
    def encerrar_api(cls) -> None:
        if cls.processo.poll() is None:
            cls.processo.terminate()
            try:
                cls.processo.wait(timeout=5)
            except subprocess.TimeoutExpired:
                cls.processo.kill()
                cls.processo.wait(timeout=5)
        if cls.processo.stdout:
            cls.processo.stdout.close()

    def requisicao(
        self,
        metodo: str,
        caminho: str,
        payload: dict | None = None,
        headers: dict[str, str] | None = None,
    ) -> tuple[int, dict[str, str], bytes]:
        corpo = json.dumps(payload).encode("utf-8") if payload is not None else None
        cabecalhos = dict(headers or {})
        if corpo is not None:
            cabecalhos["Content-Type"] = "application/json"
        request = Request(
            self.base_url + caminho,
            data=corpo,
            headers=cabecalhos,
            method=metodo,
        )
        try:
            with urlopen(request, timeout=10) as resposta:
                headers_resposta = {
                    chave.lower(): valor for chave, valor in resposta.headers.items()
                }
                return resposta.status, headers_resposta, resposta.read()
        except HTTPError as exc:
            headers_resposta = {chave.lower(): valor for chave, valor in exc.headers.items()}
            return exc.code, headers_resposta, exc.read()

    def json_request(self, metodo: str, caminho: str, payload: dict) -> tuple[int, dict]:
        status_http, _, conteudo = self.requisicao(metodo, caminho, payload)
        resposta = json.loads(conteudo)
        self.assertIsInstance(resposta, dict)
        return status_http, resposta

    def criar_empresa(self, nome: str) -> dict:
        status_http, empresa = self.json_request("POST", "/empresas", {"nome": nome})
        self.assertEqual(status_http, 201)
        return empresa

    def criar_funcionario(self, empresa_id: int, codigo: str, nome: str) -> dict:
        status_http, funcionario = self.json_request(
            "POST",
            "/funcionarios",
            {"empresa_id": empresa_id, "codigo": codigo, "nome": nome},
        )
        self.assertEqual(status_http, 201)
        return funcionario

    def criar_competencia(self, empresa_id: int) -> dict:
        status_http, competencia = self.json_request(
            "POST",
            "/competencias",
            {"empresa_id": empresa_id, "mes": 7, "ano": 2026},
        )
        self.assertEqual(status_http, 201)
        return competencia

    def criar_marcacao(self, payload: dict) -> dict:
        status_http, marcacao = self.json_request("POST", "/marcacoes", payload)
        self.assertEqual(status_http, 201)
        return marcacao

    def baixar_excel(self, competencia_id: int) -> tuple[int, dict[str, str], bytes]:
        return self.requisicao(
            "GET",
            f"/relatorios/excel?competencia_id={competencia_id}",
            headers={"Origin": "http://127.0.0.1:5500"},
        )

    def test_download_reabre_xlsx_com_abas_colunas_tipos_e_nome_corretos(self) -> None:
        empresa = self.criar_empresa("Empresa Ágil + Fórmula")
        ana = self.criar_funcionario(empresa["id"], "E001", "Ana Exportação")
        bia = self.criar_funcionario(empresa["id"], "E002", "Bia Exportação")
        caio = self.criar_funcionario(empresa["id"], "E003", "Caio Exportação")
        competencia = self.criar_competencia(empresa["id"])

        self.criar_marcacao(
            {
                "competencia_id": competencia["id"],
                "funcionario_id": ana["id"],
                "data": "2026-07-01",
                "entrada": "08:00",
                "saida_almoco": "12:00",
                "retorno_almoco": "13:00",
                "saida": "18:00",
                "status_dia": "normal",
                "origem": "manual",
                "conferido": True,
                "observacoes": "=SOMA(1;1)",
            }
        )
        self.criar_marcacao(
            {
                "competencia_id": competencia["id"],
                "funcionario_id": bia["id"],
                "data": "2026-07-01",
                "entrada": "08:30",
                "saida_almoco": "12:00",
                "retorno_almoco": "13:00",
                "saida": "17:00",
                "status_dia": "normal",
                "origem": "manual",
                "conferido": False,
            }
        )
        self.criar_marcacao(
            {
                "competencia_id": competencia["id"],
                "funcionario_id": caio["id"],
                "data": "2026-07-02",
                "entrada": "08:00",
                "status_dia": "pendente_conferencia",
                "origem": "manual",
                "conferido": False,
            }
        )

        status_http, headers, conteudo = self.baixar_excel(competencia["id"])

        self.assertEqual(status_http, 200)
        self.assertEqual(
            headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(
            headers["content-disposition"],
            'attachment; filename="on_ponto_empresa_agil_formula_2026-07.xlsx"',
        )
        self.assertIn("Content-Disposition", headers["access-control-expose-headers"])
        self.assertTrue(conteudo.startswith(b"PK"))

        workbook = load_workbook(BytesIO(conteudo), data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Resumo", "Marcações"])
        resumo = workbook["Resumo"]
        marcacoes = workbook["Marcações"]

        self.assertEqual([celula.value for celula in resumo[1]], CABECALHO_RESUMO)
        self.assertEqual([celula.value for celula in marcacoes[1]], CABECALHO_MARCACOES)
        self.assertEqual(resumo.freeze_panes, "A2")
        self.assertEqual(marcacoes.freeze_panes, "A2")
        self.assertEqual(resumo.auto_filter.ref, "A1:J4")
        self.assertEqual(marcacoes.auto_filter.ref, "A1:K4")
        self.assertTrue(resumo["A1"].fill.fgColor.rgb.endswith("16845B"))
        self.assertTrue(marcacoes["A1"].fill.fgColor.rgb.endswith("16845B"))

        linha_resumo = {
            resumo.cell(linha, 3).value: linha for linha in range(2, resumo.max_row + 1)
        }
        linha_ana = linha_resumo["Ana Exportação"]
        linha_bia = linha_resumo["Bia Exportação"]
        linha_caio = linha_resumo["Caio Exportação"]
        self.assertEqual(resumo.cell(linha_ana, 4).value, timedelta(0))
        self.assertEqual(resumo.cell(linha_ana, 5).value, timedelta(hours=1))
        self.assertEqual(resumo.cell(linha_ana, 9).value, "Conferido")
        self.assertEqual(resumo.cell(linha_bia, 4).value, timedelta(minutes=30))
        self.assertEqual(resumo.cell(linha_bia, 9).value, "Pendente")
        self.assertEqual(resumo.cell(linha_caio, 4).value, "Indisponível")
        self.assertEqual(resumo.cell(linha_caio, 5).value, "Indisponível")
        self.assertEqual(resumo.cell(linha_caio, 9).value, "Pendente")
        self.assertEqual(resumo.cell(linha_ana, 4).number_format, "[h]:mm")

        data_excel = marcacoes["A2"].value
        self.assertIsInstance(data_excel, (date, datetime))
        data_reaberta = data_excel.date() if isinstance(data_excel, datetime) else data_excel
        self.assertEqual(data_reaberta, date(2026, 7, 1))
        self.assertEqual(marcacoes["A2"].number_format, "dd/mm/yyyy")
        for coordenada, esperado in (
            ("C2", horario(8)),
            ("D2", horario(12)),
            ("E2", horario(13)),
            ("F2", horario(18)),
        ):
            self.assertIsInstance(marcacoes[coordenada].value, horario)
            self.assertEqual(marcacoes[coordenada].value, esperado)
            self.assertEqual(marcacoes[coordenada].number_format, "hh:mm")
        self.assertEqual(marcacoes["G2"].value, timedelta(hours=9))
        self.assertEqual(marcacoes["G2"].number_format, "[h]:mm")
        self.assertEqual(marcacoes["H2"].value, timedelta(hours=1))
        self.assertEqual(marcacoes["H2"].number_format, "[h]:mm")
        self.assertEqual(marcacoes["H3"].value, "-00:30")
        self.assertEqual(marcacoes["H3"].number_format, "@")
        self.assertEqual(marcacoes["G4"].value, "Indisponível")
        self.assertEqual(marcacoes["H4"].value, "Indisponível")
        self.assertEqual(marcacoes["I4"].value, "Pendente de conferência")
        self.assertEqual(marcacoes["J2"].value, "Sim")
        self.assertEqual(marcacoes["J3"].value, "Não")
        self.assertEqual(marcacoes["K2"].value, "'=SOMA(1;1)")
        self.assertEqual(marcacoes["K2"].data_type, "s")

    def test_competencia_sem_marcacoes_exporta_indisponivel_e_impressao_continua_valida(self) -> None:
        empresa = self.criar_empresa("Empresa Sem Dados")
        self.criar_funcionario(empresa["id"], "V001", "Pessoa Sem Marcações")
        competencia = self.criar_competencia(empresa["id"])

        status_http, headers, conteudo = self.baixar_excel(competencia["id"])

        self.assertEqual(status_http, 200)
        self.assertEqual(
            headers["content-disposition"],
            'attachment; filename="on_ponto_empresa_sem_dados_2026-07.xlsx"',
        )
        workbook = load_workbook(BytesIO(conteudo), data_only=False)
        self.addCleanup(workbook.close)
        resumo = workbook["Resumo"]
        marcacoes = workbook["Marcações"]
        self.assertEqual(resumo.max_row, 2)
        self.assertEqual(
            [resumo.cell(2, coluna).value for coluna in range(4, 9)],
            ["Indisponível"] * 5,
        )
        self.assertEqual(resumo["I2"].value, "Indisponível")
        self.assertEqual(marcacoes.max_row, 1)

        status_impressao, _, html = self.requisicao(
            "GET", f'/relatorios/impressao?competencia_id={competencia["id"]}'
        )
        self.assertEqual(status_impressao, 200)
        self.assertIn("Indisponível".encode("utf-8"), html)

    def test_competencia_inexistente_retorna_404_legivel(self) -> None:
        status_http, headers, conteudo = self.baixar_excel(999_999)

        self.assertEqual(status_http, 404)
        self.assertEqual(headers["content-type"], "application/json")
        self.assertEqual(json.loads(conteudo), {"detail": "Competência não encontrada."})


if __name__ == "__main__":
    unittest.main()
