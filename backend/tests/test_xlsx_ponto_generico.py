from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.importadores import xlsx_ponto_generico as adaptador
from app.importadores.xlsx_ponto_generico import (
    ErroImportacaoXlsxPontoGenerico,
    detectar,
    parse_xlsx_ponto_generico,
)


def _workbook_layout_largo() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "NUMERO DE FUNCIONARIO: 007"
    ws["B1"] = "NOME: Fulano de Tal"
    ws["C2"] = "08:00 12:00 13:00 17:00"
    ws["D2"] = "08:00 17:00"
    ws["C3"] = 1
    ws["D3"] = 2
    return wb


def _salvar_temp(wb: Workbook) -> Path:
    fd, caminho = tempfile.mkstemp(suffix=".xlsx")
    import os

    os.close(fd)
    caminho = Path(caminho)
    wb.save(caminho)
    return caminho


class DetectarTests(unittest.TestCase):
    def test_reconhece_o_layout_largo(self) -> None:
        self.assertTrue(detectar(_workbook_layout_largo()))

    def test_nao_reconhece_workbook_sem_marcador(self) -> None:
        wb = Workbook()
        wb.active["A1"] = "Qualquer coisa"
        self.assertFalse(detectar(wb))

    def test_nome_do_adaptador(self) -> None:
        self.assertEqual(adaptador.NOME_ADAPTER, "xlsx_ponto_generico")


class ParseTests(unittest.TestCase):
    def test_quatro_e_duas_batidas_sao_interpretadas(self) -> None:
        caminho = _salvar_temp(_workbook_layout_largo())
        try:
            resultado = parse_xlsx_ponto_generico(
                str(caminho), arquivo_nome="ponto.xlsx", mes=7, ano=2026, funcionarios=()
            )
        finally:
            caminho.unlink(missing_ok=True)

        self.assertEqual(resultado["total_linhas_validas"], 2)
        registros = {item["data"]: item for item in resultado["registros"]}
        dia1 = registros["2026-07-01"]
        self.assertEqual(
            dia1["interpretacao"],
            {"entrada": "08:00", "saida_intervalo": "12:00", "retorno_intervalo": "13:00", "saida": "17:00"},
        )
        self.assertEqual(dia1["status"], "nao_conferido")
        self.assertEqual(dia1["funcionario"]["codigo_origem"], "7")

        dia2 = registros["2026-07-02"]
        self.assertEqual(dia2["interpretacao"]["entrada"], "08:00")
        self.assertEqual(dia2["interpretacao"]["saida"], "17:00")
        self.assertEqual(dia2["status"], "conferir")
        self.assertFalse(dia1["fora_da_competencia"])

    def test_funcionario_cadastrado_por_codigo_e_selecionado(self) -> None:
        caminho = _salvar_temp(_workbook_layout_largo())
        try:
            resultado = parse_xlsx_ponto_generico(
                str(caminho),
                arquivo_nome="ponto.xlsx",
                mes=7,
                ano=2026,
                funcionarios=({"id": 55, "codigo": "7", "nome": "Fulano de Tal"},),
            )
        finally:
            caminho.unlink(missing_ok=True)

        dia1 = next(item for item in resultado["registros"] if item["data"] == "2026-07-01")
        self.assertTrue(dia1["funcionario"]["encontrado"])
        self.assertEqual(dia1["funcionario"]["id"], 55)
        self.assertTrue(dia1["selecionado"])

    def test_arquivo_sem_blocos_reconheciveis_gera_erro_legivel(self) -> None:
        wb = Workbook()
        wb.active["A1"] = "Qualquer coisa"
        caminho = _salvar_temp(wb)
        try:
            with self.assertRaises(ErroImportacaoXlsxPontoGenerico) as contexto:
                parse_xlsx_ponto_generico(
                    str(caminho), arquivo_nome="ponto.xlsx", mes=7, ano=2026, funcionarios=()
                )
            self.assertEqual(contexto.exception.codigo, "arquivo_vazio")
        finally:
            caminho.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
