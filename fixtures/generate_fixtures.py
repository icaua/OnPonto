from __future__ import annotations

import json
import random
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas


BASE_DIR = Path(__file__).resolve().parent
GENERATED_DIR = BASE_DIR / "generated"
PDF_DIR = GENERATED_DIR / "pdf"
IMAGES_DIR = GENERATED_DIR / "images"
SPREADSHEETS_DIR = GENERATED_DIR / "spreadsheets"
EXPECTED_DIR = GENERATED_DIR / "expected"

COMPETENCIA = "07/2026"
COMPETENCIA_MES = 7
COMPETENCIA_ANO = 2026

DIAS_SEMANA = {
    0: "segunda-feira",
    1: "terça-feira",
    2: "quarta-feira",
    3: "quinta-feira",
    4: "sexta-feira",
    5: "sábado",
    6: "domingo",
}


@dataclass(frozen=True)
class Empresa:
    nome: str
    cnpj: str


@dataclass(frozen=True)
class Funcionario:
    empresa: Empresa
    codigo: str
    nome: str


@dataclass(frozen=True)
class MarcacaoFixture:
    funcionario: Funcionario
    data: date
    entrada: str | None
    saida_almoco: str | None
    retorno_almoco: str | None
    saida: str | None
    status_dia: str
    observacao: str


def ensure_dirs() -> None:
    for pasta in [PDF_DIR, IMAGES_DIR, SPREADSHEETS_DIR, EXPECTED_DIR]:
        pasta.mkdir(parents=True, exist_ok=True)


def empresas_e_funcionarios() -> list[Funcionario]:
    mercado = Empresa("Mercado Exemplo LTDA", "12.345.678/0001-90")
    padaria = Empresa("Padaria Modelo LTDA", "23.456.789/0001-01")
    loja = Empresa("Loja Teste Comércio LTDA", "34.567.890/0001-12")

    return [
        Funcionario(mercado, "M001", "João da Silva"),
        Funcionario(mercado, "M002", "Maria Oliveira"),
        Funcionario(padaria, "P001", "Carlos Santos"),
        Funcionario(padaria, "P002", "Ana Souza"),
        Funcionario(loja, "L001", "Pedro Lima"),
        Funcionario(loja, "L002", "Fernanda Costa"),
    ]


def d(dia: int) -> date:
    return date(COMPETENCIA_ANO, COMPETENCIA_MES, dia)


def m(
    funcionario: Funcionario,
    dia: int,
    entrada: str | None,
    saida_almoco: str | None,
    retorno_almoco: str | None,
    saida: str | None,
    status_dia: str,
    observacao: str,
) -> MarcacaoFixture:
    return MarcacaoFixture(funcionario, d(dia), entrada, saida_almoco, retorno_almoco, saida, status_dia, observacao)


def gerar_marcacoes() -> list[MarcacaoFixture]:
    joao, maria, carlos, ana, pedro, fernanda = empresas_e_funcionarios()

    return [
        m(joao, 1, "08:00", "12:00", "13:00", "17:00", "normal", "Jornada normal"),
        m(joao, 2, "08:20", "12:00", "13:00", "17:00", "normal", "Pequeno atraso"),
        m(joao, 3, "08:00", "12:00", "13:00", "18:00", "normal", "Hora extra"),
        m(joao, 4, "08:00", None, None, "12:00", "normal", "Sábado com 4 horas"),
        m(joao, 5, None, None, None, None, "folga", "Domingo sem jornada"),
        m(joao, 6, None, None, None, None, "falta", "Falta sem justificativa"),
        m(joao, 7, None, None, None, None, "atestado", "Atestado médico"),
        m(joao, 8, None, None, None, None, "feriado", "Feriado municipal"),
        m(joao, 9, "08:00", "12:00", None, "17:00", "pendente", "Retorno do almoço ilegível"),
        m(joao, 10, None, "12:00", "13:00", "17:00", "pendente", "Entrada apagada na foto"),
        m(joao, 11, "08:00", None, None, "12:30", "normal", "Sábado com hora extra"),
        m(joao, 12, None, None, None, None, "folga", "Domingo sem jornada"),
        m(maria, 1, "08:05", "12:05", "13:05", "17:05", "normal", "Jornada normal"),
        m(maria, 2, "08:00", "12:00", "13:00", "16:30", "normal", "Saída antecipada"),
        m(maria, 3, "07:50", "12:00", "13:00", "17:30", "normal", "Compensação com extra"),
        m(maria, 4, "08:10", None, None, "12:00", "normal", "Sábado com pequeno atraso"),
        m(maria, 5, None, None, None, None, "folga", "Domingo sem jornada"),
        m(maria, 6, None, None, None, None, "atestado", "Atestado anexado"),
        m(maria, 7, "08:00", "12:00", "13:00", None, "pendente", "Saída ilegível"),
        m(maria, 8, None, None, None, None, "folga", "Folga compensatória"),
        m(maria, 9, "08:00", "12:00", "13:00", "17:00", "normal", "Jornada normal"),
        m(maria, 10, None, None, None, None, "falta", "Falta"),
        m(maria, 11, "08:00", None, None, "12:00", "normal", "Sábado normal"),
        m(maria, 12, None, None, None, None, "folga", "Domingo sem jornada"),
        m(carlos, 1, "08:00", "11:55", "13:05", "17:00", "normal", "Tabela desalinhada"),
        m(carlos, 2, "08:30", "12:00", "13:00", "17:00", "normal", "Chegada atrasada"),
        m(carlos, 3, "08:00", "12:00", "13:00", "18:15", "normal", "Hora extra"),
        m(carlos, 4, "08:00", None, None, "11:45", "normal", "Sábado incompleto"),
        m(carlos, 5, None, None, None, None, "folga", "Domingo"),
        m(carlos, 6, "08:00", "12:00", None, "17:00", "pendente", "Retorno não legível"),
        m(carlos, 7, None, None, None, None, "falta", "Falta"),
        m(carlos, 8, None, None, None, None, "feriado", "Feriado"),
        m(carlos, 9, "07:45", "12:00", "13:00", "17:15", "normal", "Entrada antecipada"),
        m(carlos, 10, None, None, None, None, "atestado", "Atestado"),
        m(carlos, 11, "08:00", None, None, "12:30", "normal", "Sábado extra"),
        m(carlos, 12, None, None, None, None, "folga", "Domingo"),
        m(ana, 1, "08:00", "12:00", "13:00", "17:00", "normal", "Ok"),
        m(ana, 2, "08:12", "12:00", "13:00", "17:00", "normal", "Atraso leve"),
        m(ana, 3, "08:00", "12:00", "13:00", "17:45", "normal", "Extra"),
        m(ana, 4, None, None, None, None, "folga", "Sábado de folga"),
        m(ana, 5, None, None, None, None, "folga", "Domingo"),
        m(ana, 6, None, None, None, "17:00", "pendente", "Entrada em branco"),
        m(ana, 7, None, None, None, None, "atestado", "Atestado"),
        m(ana, 8, None, None, None, None, "falta", "Falta"),
        m(ana, 9, "08:00", "12:00", "13:00", "17:00", "normal", "Ok"),
        m(ana, 10, "08:00", "12:00", "13:00", "18:30", "normal", "Plantão"),
        m(ana, 11, "08:15", None, None, "12:00", "normal", "Sábado com atraso"),
        m(ana, 12, None, None, None, None, "folga", "Domingo"),
        m(pedro, 1, "08:00", "12:00", "13:00", "17:00", "normal", "Normal"),
        m(pedro, 2, None, None, None, None, "falta", "Ausente"),
        m(pedro, 3, "08:00", "12:00", None, "17:00", "pendente", "Retorno almoço vazio"),
        m(pedro, 4, "08:00", None, None, "12:00", "normal", "Sábado normal"),
        m(pedro, 5, None, None, None, None, "folga", "Domingo"),
        m(pedro, 6, "08:45", "12:00", "13:00", "17:00", "normal", "Atraso grande"),
        m(pedro, 7, "08:00", "12:00", "13:00", "19:00", "normal", "Extra grande"),
        m(pedro, 8, None, None, None, None, "feriado", "Feriado"),
        m(pedro, 9, None, "12:00", "13:00", "17:00", "pendente", "Entrada ilegível"),
        m(pedro, 10, None, None, None, None, "atestado", "Atestado"),
        m(pedro, 11, "08:00", None, None, "11:30", "normal", "Sábado saída cedo"),
        m(pedro, 12, None, None, None, None, "folga", "Domingo"),
        m(fernanda, 1, "08:00", "12:00", "13:00", "17:00", "normal", "Normal"),
        m(fernanda, 2, "08:00", "12:00", "13:00", "17:00", "normal", "Normal"),
        m(fernanda, 3, "08:00", "12:00", "13:00", "17:20", "normal", "Extra pequena"),
        m(fernanda, 4, "08:00", None, None, "12:00", "normal", "Sábado"),
        m(fernanda, 5, None, None, None, None, "folga", "Domingo"),
        m(fernanda, 6, "08:25", "12:00", "13:00", "17:00", "normal", "Atraso"),
        m(fernanda, 7, None, None, None, None, "falta", "Falta"),
        m(fernanda, 8, None, None, None, None, "atestado", "Atestado"),
        m(fernanda, 9, "08:00", None, "13:00", "17:00", "pendente", "Saída almoço em branco"),
        m(fernanda, 10, "08:00", "12:00", "13:00", "18:00", "normal", "Extra"),
        m(fernanda, 11, None, None, None, None, "feriado", "Feriado local"),
        m(fernanda, 12, None, None, None, None, "folga", "Domingo"),
    ]


def parse_hora(valor: str | None) -> time | None:
    if not valor:
        return None
    hora, minuto = valor.split(":")
    return time(int(hora), int(minuto))


def minutos(valor: str) -> int:
    hora, minuto = valor.split(":")
    return int(hora) * 60 + int(minuto)


def fmt_minutos(total: int) -> str:
    sinal = "-" if total < 0 else ""
    total = abs(total)
    horas, minutos_restantes = divmod(total, 60)
    return f"{sinal}{horas:02d}:{minutos_restantes:02d}"


def diff(inicio: str, fim: str) -> int | None:
    valor = minutos(fim) - minutos(inicio)
    return valor if valor >= 0 else None


def jornada_prevista(data_marcacao: date) -> int:
    if data_marcacao.weekday() <= 4:
        return 8 * 60
    if data_marcacao.weekday() == 5:
        return 4 * 60
    return 0


def calcular_dia(marcacao: MarcacaoFixture) -> dict:
    prevista = jornada_prevista(marcacao.data)
    trabalhadas = 0
    atraso = 0
    extra = 0
    falta = 0
    atestado = 0
    pendente = False
    motivo_pendencia = None

    if marcacao.status_dia == "falta":
        falta = 1
    elif marcacao.status_dia == "atestado":
        atestado = 1
    elif marcacao.status_dia in {"folga", "feriado"}:
        pass
    else:
        if not marcacao.entrada or not marcacao.saida:
            pendente = True
            motivo_pendencia = "Entrada ou saída não informada."
        elif bool(marcacao.saida_almoco) != bool(marcacao.retorno_almoco):
            pendente = True
            motivo_pendencia = "Intervalo de almoço incompleto."
        else:
            if marcacao.saida_almoco and marcacao.retorno_almoco:
                manha = diff(marcacao.entrada, marcacao.saida_almoco)
                tarde = diff(marcacao.retorno_almoco, marcacao.saida)
                if manha is None or tarde is None:
                    pendente = True
                    motivo_pendencia = "Sequência de horários inválida."
                else:
                    trabalhadas = manha + tarde
            else:
                total = diff(marcacao.entrada, marcacao.saida)
                if total is None:
                    pendente = True
                    motivo_pendencia = "Sequência de horários inválida."
                else:
                    trabalhadas = total

            if not pendente:
                if trabalhadas < prevista:
                    atraso = prevista - trabalhadas
                elif trabalhadas > prevista:
                    extra = trabalhadas - prevista

        if marcacao.status_dia == "pendente":
            pendente = True
            motivo_pendencia = motivo_pendencia or "Dia marcado como pendente."

    return {
        "data": marcacao.data.isoformat(),
        "dia_semana": DIAS_SEMANA[marcacao.data.weekday()],
        "entrada": marcacao.entrada,
        "saida_almoco": marcacao.saida_almoco,
        "retorno_almoco": marcacao.retorno_almoco,
        "saida": marcacao.saida,
        "observacao": marcacao.observacao,
        "status_esperado": marcacao.status_dia,
        "jornada_prevista_minutos": prevista,
        "jornada_prevista": fmt_minutos(prevista),
        "horas_trabalhadas_minutos": trabalhadas,
        "horas_trabalhadas": fmt_minutos(trabalhadas),
        "atraso_minutos": atraso,
        "atraso": fmt_minutos(atraso),
        "extra_minutos": extra,
        "extra": fmt_minutos(extra),
        "falta": falta,
        "atestado": atestado,
        "pendente": pendente,
        "pendencia_motivo": motivo_pendencia,
    }


def agrupar_por_funcionario(marcacoes: Iterable[MarcacaoFixture]) -> dict[Funcionario, list[MarcacaoFixture]]:
    agrupado: dict[Funcionario, list[MarcacaoFixture]] = {}
    for marcacao in marcacoes:
        agrupado.setdefault(marcacao.funcionario, []).append(marcacao)
    return agrupado


def agrupar_por_empresa(marcacoes: Iterable[MarcacaoFixture]) -> dict[str, list[MarcacaoFixture]]:
    agrupado: dict[str, list[MarcacaoFixture]] = {}
    for marcacao in marcacoes:
        agrupado.setdefault(marcacao.funcionario.empresa.nome, []).append(marcacao)
    return agrupado


def resultado_esperado(marcacoes: list[MarcacaoFixture]) -> dict:
    funcionarios = []
    for funcionario, linhas in agrupar_por_funcionario(marcacoes).items():
        dias = [calcular_dia(linha) for linha in linhas]
        funcionarios.append(
            {
                "empresa": funcionario.empresa.nome,
                "cnpj": funcionario.empresa.cnpj,
                "competencia": COMPETENCIA,
                "funcionario": funcionario.nome,
                "codigo": funcionario.codigo,
                "total_atrasos_minutos": sum(dia["atraso_minutos"] for dia in dias),
                "total_atrasos": fmt_minutos(sum(dia["atraso_minutos"] for dia in dias)),
                "total_extras_minutos": sum(dia["extra_minutos"] for dia in dias),
                "total_extras": fmt_minutos(sum(dia["extra_minutos"] for dia in dias)),
                "quantidade_faltas": sum(dia["falta"] for dia in dias),
                "quantidade_atestados": sum(dia["atestado"] for dia in dias),
                "quantidade_pendencias": sum(1 for dia in dias if dia["pendente"]),
                "lista_diaria": dias,
            }
        )

    return {
        "descricao": "Resultado esperado para fixtures fictícias do On Ponto. Não contém dados reais.",
        "competencia": COMPETENCIA,
        "gerado_em": "2026-07-31T18:00:00",
        "funcionarios": funcionarios,
    }


def texto(valor: str | None) -> str:
    return valor or ""


def desenhar_tabela_pdf(canvas: Canvas, linhas: list[MarcacaoFixture], estilo: str) -> None:
    largura, altura = A4
    margem_x = 28
    y = altura - 42
    row_h = 21
    colunas = [
        ("Data", 57),
        ("Dia", 70),
        ("Entrada", 54),
        ("Saída almoço", 75),
        ("Retorno almoço", 84),
        ("Saída", 54),
        ("Observação", 166),
    ]

    funcionario = linhas[0].funcionario
    canvas.setTitle(f"Relatório ponto {funcionario.empresa.nome} {COMPETENCIA}")

    if estilo == "baguncado":
        canvas.setFillColor(colors.HexColor("#eeeeee"))
        canvas.rotate(0.15)
        canvas.rect(16, 20, largura - 32, altura - 42, fill=1, stroke=0)
        canvas.setFillColor(colors.HexColor("#222222"))

    canvas.setFont("Helvetica-Bold", 13)
    canvas.drawString(margem_x, y, "Relatório de Ponto")
    canvas.setFont("Helvetica", 9)
    y -= 18
    canvas.drawString(margem_x, y, f"Empresa: {funcionario.empresa.nome}")
    canvas.drawString(330, y, f"CNPJ: {funcionario.empresa.cnpj}")
    y -= 14
    canvas.drawString(margem_x, y, f"Funcionário: {funcionario.nome}")
    canvas.drawString(330, y, f"Código: {funcionario.codigo}")
    y -= 14
    canvas.drawString(margem_x, y, f"Competência: {COMPETENCIA}")

    if estilo == "baguncado":
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.setFont("Helvetica-Oblique", 8)
        canvas.drawString(360, y, "cópia enviada por e-mail com baixa qualidade")
        canvas.setFillColor(colors.HexColor("#222222"))

    y -= 28
    x = margem_x
    canvas.setFillColor(colors.HexColor("#e8f1ef"))
    canvas.rect(x, y, sum(largura_col for _, largura_col in colunas), row_h, fill=1, stroke=0)
    canvas.setFillColor(colors.HexColor("#222222"))
    canvas.setFont("Helvetica-Bold", 7.5)
    for titulo, largura_col in colunas:
        canvas.rect(x, y, largura_col, row_h, fill=0, stroke=1)
        canvas.drawString(x + 3, y + 7, titulo)
        x += largura_col

    canvas.setFont("Helvetica", 7.4)
    y -= row_h
    rng = random.Random(funcionario.codigo)
    for linha in linhas:
        x = margem_x
        valores = [
            linha.data.strftime("%d/%m/%Y"),
            DIAS_SEMANA[linha.data.weekday()],
            texto(linha.entrada),
            texto(linha.saida_almoco),
            texto(linha.retorno_almoco),
            texto(linha.saida),
            linha.observacao,
        ]
        if estilo == "inconsistencias" and linha.status_dia in {"pendente", "falta"}:
            canvas.setFillColor(colors.HexColor("#fff0eb"))
            canvas.rect(x, y, sum(largura_col for _, largura_col in colunas), row_h, fill=1, stroke=0)
            canvas.setFillColor(colors.HexColor("#222222"))

        for valor, (_, largura_col) in zip(valores, colunas):
            jitter_x = rng.uniform(-1.5, 1.5) if estilo == "baguncado" else 0
            jitter_y = rng.uniform(-1.0, 1.0) if estilo == "baguncado" else 0
            canvas.rect(x, y, largura_col, row_h, fill=0, stroke=1)
            canvas.drawString(x + 3 + jitter_x, y + 7 + jitter_y, str(valor)[:34])
            x += largura_col
        y -= row_h

    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#555555"))
    canvas.drawString(margem_x, 30, "Documento fictício para testes do On Ponto. Não usar como registro oficial de ponto.")


def gerar_pdf(nome_arquivo: str, marcacoes: list[MarcacaoFixture], estilo: str) -> Path:
    caminho = PDF_DIR / nome_arquivo
    canvas = Canvas(str(caminho), pagesize=A4)
    for _, linhas in agrupar_por_funcionario(marcacoes).items():
        desenhar_tabela_pdf(canvas, linhas, estilo)
        canvas.showPage()
    canvas.save()
    return caminho


def gerar_pdfs(marcacoes: list[MarcacaoFixture]) -> list[Path]:
    por_empresa = agrupar_por_empresa(marcacoes)
    return [
        gerar_pdf("relatorio_ponto_simples.pdf", por_empresa["Mercado Exemplo LTDA"], "simples"),
        gerar_pdf("relatorio_ponto_baguncado.pdf", por_empresa["Padaria Modelo LTDA"], "baguncado"),
        gerar_pdf("relatorio_ponto_inconsistencias.pdf", por_empresa["Loja Teste Comércio LTDA"], "inconsistencias"),
    ]


def fonte(preferida: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    for nome in ["arial.ttf", "DejaVuSans.ttf"]:
        try:
            return ImageFont.truetype(nome, preferida)
        except OSError:
            continue
    return ImageFont.load_default()


def gerar_imagem_ruim(marcacoes: list[MarcacaoFixture]) -> Path:
    caminho = IMAGES_DIR / "foto_whatsapp_relatorio_ruim.png"
    linhas = agrupar_por_empresa(marcacoes)["Mercado Exemplo LTDA"][:12]
    funcionario = linhas[0].funcionario

    img = Image.new("RGB", (1180, 1580), "#f8f7f0")
    draw = ImageDraw.Draw(img)
    font_titulo = fonte(34)
    font_base = fonte(21)
    font_pequena = fonte(18)

    y = 54
    draw.text((58, y), "Relatório de Ponto", fill="#1f1f1f", font=font_titulo)
    y += 52
    draw.text((58, y), f"Empresa: {funcionario.empresa.nome}", fill="#222222", font=font_base)
    y += 32
    draw.text((58, y), f"CNPJ: {funcionario.empresa.cnpj}", fill="#333333", font=font_base)
    y += 32
    draw.text((58, y), f"Funcionário: {funcionario.nome}    Código: {funcionario.codigo}", fill="#333333", font=font_base)
    y += 32
    draw.text((58, y), f"Competência: {COMPETENCIA}", fill="#333333", font=font_base)
    y += 52

    col_x = [58, 188, 336, 458, 620, 810, 930]
    headers = ["Data", "Dia", "Entrada", "Saída alm.", "Retorno", "Saída", "Obs."]
    for x, header in zip(col_x, headers):
        draw.text((x, y), header, fill="#111111", font=font_pequena)
    y += 32

    for linha in linhas:
        valores = [
            linha.data.strftime("%d/%m/%Y"),
            DIAS_SEMANA[linha.data.weekday()][:7],
            texto(linha.entrada),
            texto(linha.saida_almoco),
            texto(linha.retorno_almoco),
            texto(linha.saida),
            linha.observacao[:18],
        ]
        for x, valor in zip(col_x, valores):
            draw.text((x, y), valor, fill="#262626", font=font_pequena)
        y += 34

    rng = random.Random(202607)
    pixels = img.load()
    for _ in range(19000):
        x = rng.randrange(img.width)
        y_noise = rng.randrange(img.height)
        r, g, b = pixels[x, y_noise]
        delta = rng.randrange(-28, 24)
        pixels[x, y_noise] = (
            max(0, min(255, r + delta)),
            max(0, min(255, g + delta)),
            max(0, min(255, b + delta)),
        )

    img = img.resize((760, 1020), Image.Resampling.BILINEAR)
    img = ImageEnhance.Contrast(img).enhance(0.68)
    img = ImageEnhance.Brightness(img).enhance(0.92)
    img = img.filter(ImageFilter.GaussianBlur(radius=0.45))
    img = img.rotate(-2.4, expand=True, fillcolor="#d8d2c4")
    img.save(caminho)
    return caminho


def ajustar_larguras(ws) -> None:
    for coluna in ws.columns:
        tamanho = max(len(str(celula.value or "")) for celula in coluna)
        ws.column_dimensions[get_column_letter(coluna[0].column)].width = min(max(tamanho + 2, 12), 34)


def gerar_xlsx(marcacoes: list[MarcacaoFixture], esperado: dict) -> Path:
    caminho = SPREADSHEETS_DIR / "relatorios_ponto_ficticios.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Marcacoes"
    headers = [
        "Empresa",
        "CNPJ",
        "Competência",
        "Funcionário",
        "Código",
        "Data",
        "Dia da semana",
        "Entrada",
        "Saída Almoço",
        "Retorno Almoço",
        "Saída",
        "Observação",
    ]
    ws.append(headers)
    for linha in marcacoes:
        ws.append(
            [
                linha.funcionario.empresa.nome,
                linha.funcionario.empresa.cnpj,
                COMPETENCIA,
                linha.funcionario.nome,
                linha.funcionario.codigo,
                linha.data,
                DIAS_SEMANA[linha.data.weekday()],
                parse_hora(linha.entrada),
                parse_hora(linha.saida_almoco),
                parse_hora(linha.retorno_almoco),
                parse_hora(linha.saida),
                linha.observacao,
            ]
        )

    header_fill = PatternFill("solid", fgColor="1F6F68")
    thin = Side(style="thin", color="D7E0DF")
    for sheet in [ws]:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for row in sheet.iter_rows():
            for cell in row:
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical="top")
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=6).number_format = "yyyy-mm-dd"
            for col in [8, 9, 10, 11]:
                sheet.cell(row=row, column=col).number_format = "hh:mm"
        sheet.freeze_panes = "A2"
        ajustar_larguras(sheet)

    ws_resumo = wb.create_sheet("Resumo Esperado")
    ws_resumo.append(
        [
            "Empresa",
            "Competência",
            "Funcionário",
            "Código",
            "Atrasos",
            "Extras",
            "Faltas",
            "Atestados",
            "Pendências",
        ]
    )
    for item in esperado["funcionarios"]:
        ws_resumo.append(
            [
                item["empresa"],
                item["competencia"],
                item["funcionario"],
                item["codigo"],
                item["total_atrasos"],
                item["total_extras"],
                item["quantidade_faltas"],
                item["quantidade_atestados"],
                item["quantidade_pendencias"],
            ]
        )
    for cell in ws_resumo[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in ws_resumo.iter_rows():
        for cell in row:
            cell.border = Border(bottom=thin)
    ws_resumo.freeze_panes = "A2"
    ajustar_larguras(ws_resumo)

    wb.save(caminho)
    return caminho


def gerar_json_esperado(esperado: dict) -> Path:
    caminho = EXPECTED_DIR / "resultado_esperado.json"
    caminho.write_text(json.dumps(esperado, ensure_ascii=False, indent=2), encoding="utf-8")
    return caminho


def main() -> None:
    ensure_dirs()
    marcacoes = gerar_marcacoes()
    esperado = resultado_esperado(marcacoes)
    arquivos = []
    arquivos.extend(gerar_pdfs(marcacoes))
    arquivos.append(gerar_imagem_ruim(marcacoes))
    arquivos.append(gerar_xlsx(marcacoes, esperado))
    arquivos.append(gerar_json_esperado(esperado))

    print("Fixtures geradas:")
    for arquivo in arquivos:
        print(f"- {arquivo.relative_to(BASE_DIR.parent)}")


if __name__ == "__main__":
    main()
